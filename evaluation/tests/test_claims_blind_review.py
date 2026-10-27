"""Tests for claims/blind_review.py: seeds, commitments and the blind workbook leak audit.

No network, no LLM: openpyxl round trips over tmp_path fixtures.
"""

from __future__ import annotations

import csv
import dataclasses
import hashlib
import json
import os
import random
import re
import types
import zipfile
from collections import Counter
from pathlib import Path

import blind_review as br
import export_review_sheet as ex
import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from build_hss_set import doi_slug
from common import read_jsonl, write_json

FIXED = br.BlindSeeds(salt="00" * 16, claims_id_seed=11, claims_row_seeds=(21, 22, 23),
                      screening_sample_seed=31, screening_id_seed=32,
                      screening_row_seeds=(41, 42, 43))

PROPS = ("creator", "title", "subject", "description", "keywords", "category",
         "lastModifiedBy")


def _col(header, name):
    return openpyxl.utils.get_column_letter(list(header).index(name) + 1)


def _dv(ws, col_letter, choices, n):
    dv = DataValidation(type="list", formula1=f'"{",".join(choices)}"', allow_blank=True,
                        showErrorMessage=True)
    dv.add(f"{col_letter}2:{col_letter}{n + 1}")
    ws.add_data_validation(dv)


def _finish(wb, path, leak):
    for prop in PROPS:
        setattr(wb.properties, prop, "")
    if leak in PROPS:
        setattr(wb.properties, leak, "ScholarRAG tool")
    if leak == "defined_name":
        wb.defined_names["x"] = DefinedName("x", attr_text="Guide!$A$1")
    wb.save(path)
    if leak and leak.startswith("part_"):
        with zipfile.ZipFile(path, "a") as zf:
            zf.writestr(leak[len("part_"):], b"x")
    return path


def _clean_claims_wb(path, *, leak=None, claim="A claim.", n=2, row_order=None):
    path = Path(path)
    wb = openpyxl.Workbook()
    guide = wb.active
    guide.title = "Guide"
    guide["A1"] = "Read each claim and the source text."
    if leak == "em_dash_in_guide":
        guide["A2"] = "Do this — then that."
    if leak == "guide_word_verifier":
        guide["A2"] = "The verifier said so."
    if leak == "hidden_sheet":
        wb.create_sheet("Hidden").sheet_state = "hidden"
    ws = wb.create_sheet("Claims")
    header = list(br.BLIND_CLAIMS_HEADER)
    if leak == "header_verifier_status":
        header.append("verifier_status")
    ws.append(header)
    order = row_order if row_order is not None else list(range(n))
    for i in order:
        iid = f"BC-{i + 1:03d}"
        if leak == "bad_id_R-001" and i == 0:
            iid = "R-001"
        vals = {"item_id": iid, "claim": claim, "cited_title": "Title", "cited_doi": "10.1/x",
                "cited_first_author_year": "Smith 2020", "stored_text": "stored"}
        if leak == "real_item_id_in_note" and i == 0:
            vals["human_note"] = "see hss-econ-03"
        ws.append([vals.get(h, "") for h in header])
    if leak == "hidden_column":
        ws.column_dimensions["B"].hidden = True
    if leak == "hidden_row":
        ws.row_dimensions[2].hidden = True
    if leak == "comment":
        ws["A2"].comment = Comment("tool says", "x")
    if leak == "formula":
        ws["K2"] = "=1+1"
    _dv(ws, _col(br.BLIND_CLAIMS_HEADER, "human_verdict"), ex.HUMAN_VERDICT_CHOICES, n)
    _dv(ws, _col(br.BLIND_CLAIMS_HEADER, "confidence"), br.CONFIDENCE_CHOICES, n)
    if leak == "dropdown_extra_choice":
        _dv(ws, "I", ("verified", "constructed"), n)
    if leak == "dropdown_accepts_any_value":
        open_dv = DataValidation(type="list", formula1='"high,medium,low"', allow_blank=True,
                                 showErrorMessage=False)
        open_dv.add("I2:I3")
        ws.add_data_validation(open_dv)
    src = wb.create_sheet("Sources")
    sh = list(br.SOURCES_HEADER)
    if leak == "sources_header_extra_column":
        sh.append("extra")
    src.append(sh)
    src.append(["BC-001", 1, "te xt" if leak == "nbsp_in_sources" else "text"])
    att = wb.create_sheet("Attestation")
    att.append(list(br.ATTESTATION_HEADER))
    att.append(["Did you see tool output?", ""])
    _dv(att, "B", ("yes", "no"), 1)
    if leak == "extra_sheet_Model_labels":
        wb.create_sheet("Model_labels")
    return _finish(wb, path, leak)


def _clean_screening_wb(path, *, guide_extra="", criteria=("I1", "E1", "TOPIC")):
    path = Path(path)
    wb = openpyxl.Workbook()
    guide = wb.active
    guide.title = "Guide"
    guide["A1"] = "Decide on each record."
    if guide_extra:
        guide["A2"] = guide_extra
    ws = wb.create_sheet("Screening")
    ws.append(list(br.BLIND_SCREENING_HEADER))
    for i in (1, 2):
        vals = {"record_id": f"BS-{i:03d}", "title": "T", "abstract": "Abstract",
                "venue": "Venue", "year": 2020}
        ws.append([vals.get(h, "") for h in br.BLIND_SCREENING_HEADER])
    _dv(ws, _col(br.BLIND_SCREENING_HEADER, "human_decision"), br.SCREENING_DECISIONS, 2)
    _dv(ws, _col(br.BLIND_SCREENING_HEADER, "human_criterion"), criteria, 2)
    att = wb.create_sheet("Attestation")
    att.append(list(br.ATTESTATION_HEADER))
    att.append(["Did you see tool output?", ""])
    _dv(att, "B", ("yes", "no"), 1)
    return _finish(wb, path, None)


def _key(path, sets):
    with Path(path).open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["opaque_id", "real_item_id", "set",
                                           "construction_category", "expected_label",
                                           "workbook_id"])
        w.writeheader()
        for i, s in enumerate(sets):
            w.writerow({"opaque_id": f"BC-{i + 1:03d}", "real_item_id": f"real-x-{i:02d}",
                        "set": s, "construction_category": "", "expected_label": "",
                        "workbook_id": "w"})
    return Path(path)


def test_seed_commitment_is_sha256_of_seed_colon_salt():
    assert br.seed_commitment(7, "ab") == hashlib.sha256(b"7:ab").hexdigest()


def test_blind_workbook_id_is_salted():
    a = br.blind_workbook_id(["x", "y"], 5, "s1")
    assert len(a) == 16 and a != br.blind_workbook_id(["x", "y"], 5, "s2")
    assert a != ex.compute_workbook_id(["x", "y"], 5)


def test_seeds_file_round_trips_and_refuses_overwrite(tmp_path):
    s = br.new_seeds()
    p = tmp_path / "s.json"
    br.write_seeds_file(p, s)
    assert br.load_seeds(p) == s
    with pytest.raises(FileExistsError):
        br.write_seeds_file(p, s)


def test_new_seeds_are_in_range_and_salted():
    s = br.new_seeds()
    assert len(s.salt) == 32
    for v in (s.claims_id_seed, s.screening_sample_seed, s.screening_id_seed,
              *s.claims_row_seeds, *s.screening_row_seeds):
        assert 1 <= v <= 2**31 - 1


def test_row_seed_per_copy_and_rejects_unknown_copy():
    assert [br.row_seed(FIXED, "screening", c) for c in ("1", "2", "3")] == [41, 42, 43]
    with pytest.raises(ValueError):
        br.row_seed(FIXED, "claims", "4")
    with pytest.raises(ValueError):
        br.id_seed(FIXED, "claims", "4")


@pytest.mark.parametrize("kind", ["claims", "screening"])
@pytest.mark.parametrize("copy", ["1", "2", "3"])
def test_every_copy_uses_distinct_id_and_row_seeds(kind, copy):
    assert br.id_seed(FIXED, kind, copy) != br.row_seed(FIXED, kind, copy)


def test_builder_refuses_equal_id_and_row_seeds(tmp_path):
    bad = dataclasses.replace(FIXED, claims_row_seeds=(11, 22, 23))
    items_path, cache = _write_items(tmp_path, [_constructed(i) for i in range(1, 5)])
    with pytest.raises(ex.ExportError, match="row order"):
        br.build_blind_claims_workbook([(items_path, cache)], year_sources=[], seeds=bad,
                                       copy="1", sources_max_chars=30_000,
                                       max_cell_chars=32_000)
    bad_s = dataclasses.replace(FIXED, screening_row_seeds=(32, 42, 43))
    rec = tmp_path / "rec.json"
    rec.write_text(json.dumps(_screening_record_dict()), encoding="utf-8")
    excluded = tmp_path / "excluded.txt"
    excluded.write_text("title:paper excluded zero\n", encoding="utf-8")
    with pytest.raises(ex.ExportError, match="row order"):
        br.build_blind_screening_workbook(
            rec, expect_sha256=_sha(rec), exclude_keys_path=excluded, seeds=bad_s, copy="1",
            max_cell_chars=32_000, sizes=BUILD_SIZES)


def test_audit_flags_row_order_that_repeats_id_order(tmp_path):
    path = _clean_claims_wb(tmp_path / "a.xlsx", n=6)
    key = _key(tmp_path / "a.csv", ["constructed"] * 3 + ["real"] * 3)
    found = br.audit_blind_workbook(path, kind="claims", key_path=key)
    assert any("repeats the id order" in f for f in found)
    path2 = _clean_claims_wb(tmp_path / "b.xlsx", n=6, row_order=[1, 0, 3, 2, 5, 4])
    key2 = _key(tmp_path / "b.csv",
                ["constructed", "constructed", "real", "constructed", "real", "real"])
    found2 = br.audit_blind_workbook(path2, kind="claims", key_path=key2)
    assert any("repeats the id order" in f for f in found2)
    path3 = _clean_claims_wb(tmp_path / "c.xlsx", n=6)
    key3 = _key(tmp_path / "c.csv",
                ["constructed", "real", "constructed", "real", "real", "constructed"])
    assert br.audit_blind_workbook(path3, kind="claims", key_path=key3) == []


def test_min_set_changes_scales_with_n():
    assert br.min_set_changes({"a": 20, "b": 20}) == 10
    assert br.min_set_changes({"constructed": 60, "real": 26}) == 19
    assert br.min_set_changes(
        {"exclude": 80, "include": 30, "judge_demoted": 20, "other_needs_review": 10}) == 43
    assert br.min_set_changes({"constructed": 3, "real": 3}) is None


def test_audit_clean_workbook_has_no_findings(tmp_path):
    assert br.audit_blind_workbook(_clean_claims_wb(tmp_path / "c.xlsx"), kind="claims") == []


@pytest.mark.parametrize("leak", [
    "hidden_column", "hidden_row", "hidden_sheet", "comment",
    "formula", "defined_name", "creator", "title", "subject", "description", "keywords",
    "category", "lastModifiedBy", "part_xl/comments1.xml",
    "part_xl/externalLinks/externalLink1.xml",
    "part_xl/vbaProject.bin", "part_customXml/item1.xml", "part_xl/embeddings/oleObject1.bin",
    "extra_sheet_Model_labels", "header_verifier_status", "sources_header_extra_column",
    "guide_word_verifier", "em_dash_in_guide", "real_item_id_in_note", "bad_id_R-001",
    "dropdown_extra_choice", "dropdown_accepts_any_value", "nbsp_in_sources"])
def test_audit_reports_each_leak(tmp_path, leak):
    path = _clean_claims_wb(tmp_path / "c.xlsx", leak=leak)
    assert br.audit_blind_workbook(path, kind="claims") != []


def test_audit_flags_unsearchable_sources_character(tmp_path):
    path = _clean_claims_wb(tmp_path / "c.xlsx", leak="nbsp_in_sources")
    findings = br.audit_blind_workbook(path, kind="claims")
    assert any("Sources" in f and "U+00A0" in f and "typed search" in f for f in findings)


def test_audit_flags_scholarrag_in_file_name(tmp_path):
    path = _clean_claims_wb(tmp_path / "ScholarRAG-c.xlsx")
    assert br.audit_blind_workbook(path, kind="claims") != []


def test_audit_ignores_forbidden_words_inside_free_text(tmp_path):
    path = _clean_claims_wb(tmp_path / "c.xlsx", claim="The verifier constructed a scale.")
    assert br.audit_blind_workbook(path, kind="claims") == []


def test_audit_with_key_flags_grouped_order(tmp_path):
    path = _clean_claims_wb(tmp_path / "c.xlsx", n=40)
    key = _key(tmp_path / "k.csv", ["constructed"] * 20 + ["real"] * 20)
    assert any("order" in f for f in br.audit_blind_workbook(path, kind="claims", key_path=key))


def test_audit_with_key_passes_fixed_shuffle(tmp_path):
    sets = ["constructed"] * 20 + ["real"] * 20
    random.Random(7).shuffle(sets)
    path = _clean_claims_wb(tmp_path / "c.xlsx", n=40, row_order=list(range(40))[::-1])
    key = _key(tmp_path / "k.csv", sets)
    assert br.audit_blind_workbook(path, kind="claims", key_path=key) == []


def test_audit_with_key_flags_id_mismatch(tmp_path):
    path = _clean_claims_wb(tmp_path / "c.xlsx", n=6)
    key = _key(tmp_path / "k.csv", ["constructed"] * 3 + ["real"] * 2)
    assert br.audit_blind_workbook(path, kind="claims", key_path=key) != []


def test_audit_skips_order_check_below_30_rows(tmp_path):
    path = _clean_claims_wb(tmp_path / "c.xlsx", n=6)
    key = _key(tmp_path / "k.csv",
               ["constructed", "real", "constructed", "real", "real", "constructed"])
    assert br.audit_blind_workbook(path, kind="claims", key_path=key) == []


def test_audit_screening_flags_uppercase_status_token(tmp_path):
    path = _clean_screening_wb(tmp_path / "s.xlsx", guide_extra="EXCLUDE")
    assert br.audit_blind_workbook(path, kind="screening") != []


def test_audit_screening_accepts_criterion_dropdown(tmp_path):
    path = _clean_screening_wb(tmp_path / "s.xlsx", criteria=("I1", "I2", "E1", "TOPIC"))
    assert br.audit_blind_workbook(path, kind="screening") == []


# --------------------------------------------------------------------------------------
# Task 2: blind claims workbook and the --blind dispatch
# --------------------------------------------------------------------------------------

EVALUATION_DIR = Path(__file__).resolve().parent.parent
CLAIMS_DIR = EVALUATION_DIR / "claims"
SHARED_DOI = "10.1/shared"


def _make_cache(cache_dir, doi, chunks):
    cache_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "doi": doi, "title": "Cached paper", "authors": ["Author One"],
        "source_url": "https://example.org/SENTINELURL", "n_chars": sum(map(len, chunks)),
        "chunks": [{"section": "results", "text": c} for c in chunks],
    }
    write_json(cache_dir / f"{doi_slug(doi)}.json", payload)


def _constructed(i, *, rule="verbatim", chunk=True, doi=None):
    doi = doi or f"10.1/c{i}"
    return {
        "item_id": f"hss-{rule}-{i:02d}", "rule": "over_specified" if rule == "verbatim" else rule,
        "expected": ["needs_nuance"], "claim": f"Constructed claim number {i} about learning.",
        "original_sentence": "SENTINELORIG", "alteration": "SENTINELALT",
        "source_doi": doi, "chunk_doi": doi if chunk else None, "chunk_index": 0,
        "title": f"Constructed paper {i}", "authors": ["A. One", "B. Two"],
        "built_at": "SENTINELBUILT",
    }


def _real(i, doi=None):
    doi = doi or f"10.2/r{i}"
    return {
        "item_id": f"real-test-{i:02d}", "rule": "real", "expected": None,
        "claim": f"Real claim number {i} about teachers.", "cited_doi": doi,
        "cited_title": f"Real cited paper {i}", "cited_year": "2018", "chunk_doi": doi,
        "source_doi": "10.9/review", "title": "x", "authors": ["C. Three"],
        "review_title": "SENTINELREVIEW", "context_before": "SENTINELCTX",
        "citation_text": "SENTINELCITE", "resolution": "SENTINELRES",
    }


def _write_items(tmp_path, items, caches=None, name="items.jsonl"):
    cache = tmp_path / "cache"
    cache.mkdir(exist_ok=True)
    for it in items:
        doi = it.get("chunk_doi")
        if doi:
            _make_cache(cache, doi, (caches or {}).get(doi, [f"Text of {doi}. Learning."]))
    path = tmp_path / name
    path.write_text("\n".join(json.dumps(it) for it in items) + "\n", encoding="utf-8")
    return path, cache


def _argv(tmp_path, items_path, cache, copy, tag, seeds_path, *extra):
    return ["--items", str(items_path), "--cache-dir", str(cache), "--year-sources",
            str(tmp_path / "years.json"), "--seeds-file", str(seeds_path), "--copy", copy,
            "--out", str(tmp_path / f"{tag}.xlsx"), "--key-out", str(tmp_path / f"{tag}.csv"),
            *extra]


def _seeds(tmp_path):
    path = tmp_path / "seeds.json"
    if not path.exists():
        br.write_seeds_file(path, FIXED)
    (tmp_path / "years.json").write_text("[]", encoding="utf-8")
    return path


def _sheet_dicts(path, sheet):
    ws = openpyxl.load_workbook(path)[sheet]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, [c.value for c in row], strict=True))
            for row in ws.iter_rows(min_row=2) if any(c.value is not None for c in row)]


def _column_a(path, sheet):
    ws = openpyxl.load_workbook(path)[sheet]
    return [str(r[0].value) for r in ws.iter_rows() if r[0].value is not None]


def _zip_texts(path):
    with zipfile.ZipFile(path) as zf:
        return [(n, zf.read(n).decode("utf-8", errors="replace")) for n in zf.namelist()]


def _cells(rng):
    return [f"{openpyxl.utils.get_column_letter(c)}{r}"
            for r in range(rng.min_row, rng.max_row + 1)
            for c in range(rng.min_col, rng.max_col + 1)]


@pytest.fixture
def built(tmp_path):
    items = [_constructed(i) for i in range(1, 10)]
    items.append(_constructed(10, rule="no_full_text", chunk=False))
    items += [_real(i) for i in range(1, 11)]
    items_path, cache = _write_items(tmp_path, items)
    seeds_path = _seeds(tmp_path)
    argv = _argv(tmp_path, items_path, cache, "1", "copy1", seeds_path)
    assert br.main(["claims", *argv]) == 0
    return types.SimpleNamespace(
        out=tmp_path / "copy1.xlsx", key=tmp_path / "copy1.csv", n_items=len(items),
        workbook_id=br.read_guide_workbook_id(tmp_path / "copy1.xlsx"), tmp=tmp_path,
        items_path=items_path, cache=cache, seeds=seeds_path)


def test_claims_header_and_sheets_are_the_whitelist(built):
    wb = openpyxl.load_workbook(built.out)
    assert tuple(wb.sheetnames) == br.CLAIMS_SHEETS
    assert tuple(c.value for c in wb["Claims"][1]) == br.BLIND_CLAIMS_HEADER
    assert tuple(c.value for c in wb["Sources"][1]) == br.SOURCES_HEADER


def test_no_sentinel_item_id_or_category_in_any_xml_part(built):
    for name, text in _zip_texts(built.out):
        assert "SENTINEL" not in text, name
        assert not ex.REAL_ITEM_ID_PATTERN.search(text), name
        assert "over_specified" not in text and "constructed" not in text, name


def test_all_items_kept_and_no_full_text_row_carries_marker_and_no_sources(built):
    rows = _sheet_dicts(built.out, "Claims")
    assert len(rows) == built.n_items == 20
    marker_ids = {r["item_id"] for r in rows
                  if r["stored_text"] == ex.NO_FULL_TEXT_EVIDENCE_MARKER}
    assert len(marker_ids) == 1
    assert marker_ids.isdisjoint({r["item_id"] for r in _sheet_dicts(built.out, "Sources")})


def test_same_doi_cited_by_marker_item_and_text_item(tmp_path):
    items = [_constructed(1, rule="no_full_text", chunk=False, doi=SHARED_DOI),
             _constructed(2, doi=SHARED_DOI)]
    items_path, cache = _write_items(tmp_path, items)
    seeds_path = _seeds(tmp_path)
    assert br.main(["claims", *_argv(tmp_path, items_path, cache, "1", "o", seeds_path)]) == 0
    out = tmp_path / "o.xlsx"
    rows, src = _sheet_dicts(out, "Claims"), _sheet_dicts(out, "Sources")
    marker = [r["item_id"] for r in rows if r["stored_text"] == ex.NO_FULL_TEXT_EVIDENCE_MARKER]
    other = [r["item_id"] for r in rows if r["stored_text"] == br.STORED_TEXT_SEE_SOURCES]
    assert len(marker) == 1 and len(other) == 1
    assert sum(s["item_id"] == marker[0] for s in src) == 0
    assert sum(s["item_id"] == other[0] for s in src) >= 1


def test_control_character_is_stripped_and_long_text_split(tmp_path):
    items_path, cache = _write_items(
        tmp_path, [_constructed(1)], caches={"10.1/c1": ["a\x0cb" * 40_000]})
    seeds_path = _seeds(tmp_path)
    args = _argv(tmp_path, items_path, cache, "1", "o", seeds_path,
                 "--sources-max-chars", "30000")
    assert br.main(["claims", *args]) == 0
    src = _sheet_dicts(tmp_path / "o.xlsx", "Sources")
    assert [s["part"] for s in src] == list(range(1, len(src) + 1)) and len(src) >= 3
    limit = br._SOURCES_MAX_LINES * br._SOURCES_WIDTH
    assert all("\x0c" not in s["text"] and len(s["text"]) <= limit for s in src)
    assert sum(len(s["text"]) for s in src) == 80_000


def test_missing_cached_text_refuses(tmp_path):
    items_path = tmp_path / "items.jsonl"
    items_path.write_text(json.dumps(_constructed(1, doi="10.1/absent")) + "\n",
                          encoding="utf-8")
    cache = tmp_path / "emptycache"
    cache.mkdir()
    with pytest.raises(ex.ExportError, match="no cached text"):
        br.build_blind_claims_workbook([(items_path, cache)], year_sources=[], seeds=FIXED,
                                       copy="1", sources_max_chars=30_000,
                                       max_cell_chars=32_000)


def test_unknown_select_id_refuses(tmp_path):
    items_path, cache = _write_items(tmp_path, [_constructed(1)])
    with pytest.raises(ex.ExportError, match="unknown select id"):
        br.build_blind_claims_workbook([(items_path, cache)], year_sources=[], seeds=FIXED,
                                       copy="1", sources_max_chars=30_000,
                                       max_cell_chars=32_000, select_ids=["nope-01"])


def test_build_via_select_ids_file_uses_the_claims_id_seed_and_subset(built):
    chosen = ["hss-verbatim-01", "hss-verbatim-02", "real-test-01", "real-test-02"]
    ids_file = built.tmp / "subset_ids.txt"
    ids_file.write_text("\n".join(chosen) + "\n\n", encoding="utf-8")
    args = _argv(built.tmp, built.items_path, built.cache, "1", "subset", built.seeds,
                 "--select-ids", str(ids_file), "--expect-items", "4")
    assert br.main(["claims", *args]) == 0
    rows = _sheet_dicts(built.tmp / "subset.xlsx", "Claims")
    assert sorted(r["item_id"] for r in rows) == ["BC-001", "BC-002", "BC-003", "BC-004"]
    key = list(csv.DictReader((built.tmp / "subset.csv").open(encoding="utf-8", newline="")))
    assert sorted(r["real_item_id"] for r in key) == sorted(chosen)
    # the id order comes from claims_id_seed
    specs = ex.resolve_item_specs([built.items_path], [built.cache])
    items = [it for it in ex.read_items(specs) if it["item_id"] in set(chosen)]
    expected = list(items)
    random.Random(br.id_seed(FIXED, "claims", "1")).shuffle(expected)
    by_open = {r["opaque_id"]: r["real_item_id"] for r in key}
    assert [by_open[f"BC-{i:03d}"] for i in range(1, 5)] == [it["item_id"] for it in expected]
    assert br.read_guide_workbook_id(built.tmp / "subset.xlsx") != built.workbook_id


def test_select_ids_unknown_id_in_file_exits_1(built, capsys):
    ids_file = built.tmp / "bad_ids.txt"
    ids_file.write_text("real-test-01\nnope-99\n", encoding="utf-8")
    args = _argv(built.tmp, built.items_path, built.cache, "1", "bad", built.seeds,
                 "--select-ids", str(ids_file))
    assert br.main(["claims", *args]) == 1
    assert "unknown select id: nope-99" in capsys.readouterr().out


def test_first_author_year_matches_build_sheet_v3_rule():
    assert br.first_author_year(["Ahmet Cihat Yavuz"], 2020) == "Yavuz, 2020"
    assert br.first_author_year(["A One", "B Two"], None) == "One & Two, n.d."
    assert br.first_author_year(["A One", "B Two", "C Three"], "2018") == "One et al., 2018"
    assert br.first_author_year(None, 2020) == "Unknown, 2020"


def test_year_index_prefers_source_list_then_cited_year(tmp_path):
    src = tmp_path / "s.json"
    src.write_text(json.dumps([{"doi": "10.1/c1", "year": 2020}]), encoding="utf-8")
    idx = br.load_year_index([src])
    assert idx == {"10.1/c1": 2020}
    row = br.blind_claim_row(_constructed(1), "BC-001", year_index=idx, max_cell_chars=32_000)
    assert row["cited_first_author_year"] == "One & Two, 2020"
    row = br.blind_claim_row(_real(1), "BC-002", year_index=idx, max_cell_chars=32_000)
    assert row["cited_first_author_year"] == "Three, 2018"


def test_copies_share_ids_and_key_but_differ_in_row_order(built):
    args = _argv(built.tmp, built.items_path, built.cache, "2", "copy2", built.seeds)
    assert br.main(["claims", *args]) == 0
    rows1 = _sheet_dicts(built.out, "Claims")
    rows2 = _sheet_dicts(built.tmp / "copy2.xlsx", "Claims")
    assert ({r["item_id"]: r["claim"] for r in rows1}
            == {r["item_id"]: r["claim"] for r in rows2})
    assert built.key.read_bytes() == (built.tmp / "copy2.csv").read_bytes()
    assert [r["item_id"] for r in rows1] != [r["item_id"] for r in rows2]


def test_guide_carries_rubric_verbatim_and_workbook_id(built):
    guide = "\n".join(_column_a(built.out, "Guide"))
    for d in ex.STATUS_DEFINITIONS_V3:
        assert d in guide
    assert br.read_guide_workbook_id(built.out) == built.workbook_id
    assert "There are 20 rows." in guide


def test_attestation_rows_and_yes_no_dropdowns(built):
    ws = openpyxl.load_workbook(built.out)["Attestation"]
    assert tuple(c.value for c in ws[1]) == br.ATTESTATION_HEADER
    assert ([ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
            == [q for q, _ in br.ATTESTATION_QUESTIONS])
    want = {f"B{i + 2}" for i, (_, k) in enumerate(br.ATTESTATION_QUESTIONS) if k == "yes_no"}
    got = {c for dv in ws.data_validations.dataValidation if dv.formula1 == '"yes,no"'
           for rng in dv.sqref.ranges for c in _cells(rng)}
    assert got == want


def test_layout_widths_heights_and_panes(built):
    wb = openpyxl.load_workbook(built.out)
    ws = wb["Claims"]
    header = [c.value for c in ws[1]]

    def width(name):
        return ws.column_dimensions[_col(header, name)].width

    assert ws.freeze_panes == "C2" and width("claim") == 50
    assert width("stored_text") == 30 and ws.row_dimensions[2].height == 45
    assert width("human_quote") == 30 and width("human_note") == 30
    assert wb["Sources"].row_dimensions[2].height == 409
    assert wb["Sources"].freeze_panes == "A2"
    assert wb["Guide"].column_dimensions["A"].width == 120
    assert wb["Attestation"].column_dimensions["A"].width == 70
    verdict = {c for dv in ws.data_validations.dataValidation
               if dv.formula1 == '"' + ",".join(ex.HUMAN_VERDICT_CHOICES) + '"'
               for rng in dv.sqref.ranges for c in _cells(rng)}
    assert verdict == {f"C{r}" for r in range(2, 22)}


def test_core_properties_are_empty(built):
    props = openpyxl.load_workbook(built.out).properties
    for prop in br._PROPERTY_NAMES:
        assert getattr(props, prop) in (None, ""), prop


def test_built_workbook_passes_audit_with_key(built):
    assert br.audit_blind_workbook(built.out, kind="claims", key_path=built.key) == []


def test_key_maps_every_id_and_records_workbook_id(built):
    rows = list(csv.DictReader(built.key.open(encoding="utf-8", newline="")))
    assert {r["opaque_id"] for r in rows} == {
        r["item_id"] for r in _sheet_dicts(built.out, "Claims")}
    assert {r["workbook_id"] for r in rows} == {built.workbook_id}


def test_expect_items_mismatch_exits_2(built):
    args = _argv(built.tmp, built.items_path, built.cache, "3", "copy3", built.seeds,
                 "--expect-items", "19")
    assert br.main(["claims", *args]) == 2


def test_export_review_sheet_blind_dispatch_matches_blind_review(built):
    a = _argv(built.tmp, built.items_path, built.cache, "1", "a", built.seeds)
    b = _argv(built.tmp, built.items_path, built.cache, "1", "b", built.seeds)
    assert ex.main(["--blind", *a]) == 0 and br.main(["claims", *b]) == 0
    for sheet in ("Claims", "Sources"):
        assert (_sheet_dicts(built.tmp / "a.xlsx", sheet)
                == _sheet_dicts(built.tmp / "b.xlsx", sheet))


@pytest.mark.parametrize("flag", ["--annotation", "--results-dir", "--screening-record",
                                  "--delivered-evidence"])
def test_blind_refuses_tool_flags(built, flag):
    a = _argv(built.tmp, built.items_path, built.cache, "1", "a", built.seeds)
    with pytest.raises(SystemExit) as e:
        ex.main(["--blind", flag, "x", *a])
    assert e.value.code == 2


def test_committed_item_files_yield_86_items_with_8_markers():
    items = (read_jsonl(CLAIMS_DIR / "hss_test_v3_claims.jsonl")
             + read_jsonl(CLAIMS_DIR / "real_claims_test.jsonl"))
    idx = br.load_year_index([CLAIMS_DIR / "hss_test_v3_sources.json"])
    rows = [br.blind_claim_row(it, f"BC-{i:03d}", year_index=idx, max_cell_chars=32_000)
            for i, it in enumerate(items, 1)]
    assert len(rows) == 86
    assert sum(r["stored_text"] == ex.NO_FULL_TEXT_EVIDENCE_MARKER for r in rows) == 8
    assert all(not r["cited_first_author_year"].endswith("n.d.")
               for r, it in zip(rows, items, strict=True) if it["rule"] == "real")


# --------------------------------------------------------------------------------------
# Citation-free claim text
# --------------------------------------------------------------------------------------

CITATION_CHECK = re.compile(
    r"\([^()]*\b(?:19|20)\d{2}[a-z]?\b[^()]*\)|\w\s*\((?:19|20)\d{2}[a-z]?\)|et al\.")

GOLDEN_CLAIMS = {
    "real-test-01": "In their study, they found a positive relation between teacher agency "
                    "and openness.",
    "real-test-02": "Four teacher profiles (Carefree, Pragmatic, Uncertain and Worried) were "
                    "identified by them based on how open and positive the teachers feel "
                    "about innovations.",
    "real-test-05": "In their study, teachers reported to understand what was expected from "
                    "them in their traditional directive role, but they did not fully "
                    "understand what was expected from them in their new facilitative role.",
    "real-test-07": "The author adds a focus on media literacy, pointing to the importance of "
                    "critical thinking and personal engagement in interpreting and deriving "
                    "significance from media messages.",
    "real-test-11": "They found constructivist beliefs about mathematics teaching to be "
                    "moderately correlated with noticing (r = 0.35).",
    "real-test-15": "A Palestinian study on the use of Facebook by students in an "
                    "undergraduate educational technology course revealed strong acceptance "
                    "of this delivery modality as an effective way to learn.",
    "real-test-19": "There has been at least one formal mediation analysis that showed how "
                    "improvements in caregiver book-sharing competence associated with "
                    "training accounted for much of the improvement in child language.",
    "real-test-22": "They conducted a comprehensive overview of factors determining the level "
                    "of cognitive load and identified age as one of the causal factors.",
    "hss-over_specified-07": "In the present study, PQ Method was employed to produce the "
                             "initial by-person correlation matrix, using eye-tracking "
                             "software.",
}


def _committed_items():
    return (read_jsonl(CLAIMS_DIR / "hss_test_v3_claims.jsonl")
            + read_jsonl(CLAIMS_DIR / "real_claims_test.jsonl"))


def _stored_claim(item_id):
    return next(it["claim"] for it in _committed_items() if it["item_id"] == item_id)


@pytest.mark.parametrize("item_id", sorted(GOLDEN_CLAIMS))
def test_neutralise_citations_golden(item_id):
    assert br.neutralise_citations(_stored_claim(item_id)) == GOLDEN_CLAIMS[item_id]


def test_neutralise_citations_golden_altered_08():
    stored = _stored_claim("hss-altered-08")
    removed = " (Acheson & Nelson, 2020; De Costa et al., 2020; Hall, 2017)"
    assert removed in stored
    out = br.neutralise_citations(stored)
    assert out == stored.replace(removed, "")
    assert out.endswith("on various psychological aspects.")


def _all_blind_rows():
    items = _committed_items()
    idx = br.load_year_index([CLAIMS_DIR / "hss_test_v3_sources.json"])
    return items, [br.blind_claim_row(it, f"BC-{i:03d}", year_index=idx, max_cell_chars=32_000)
                   for i, it in enumerate(items, 1)]


def test_displayed_claims_carry_no_in_text_citation_in_either_set():
    items, rows = _all_blind_rows()
    real = [(it, r) for it, r in zip(items, rows, strict=True) if it["rule"] == "real"]
    cons = [(it, r) for it, r in zip(items, rows, strict=True) if it["rule"] != "real"]
    assert (len(real), len(cons)) == (26, 60)
    assert sum(bool(CITATION_CHECK.search(r["claim"])) for _, r in real) == 0
    assert sum(bool(CITATION_CHECK.search(r["claim"])) for _, r in cons) == 0
    assert sum(bool(CITATION_CHECK.search(it["claim"])) for it, _ in real) == 26
    assert sum(bool(CITATION_CHECK.search(it["claim"])) for it, _ in cons) == 6


def test_citation_rule_changes_only_citations():
    items, rows = _all_blind_rows()
    changed = {"real": 0, "constructed": 0}
    for it, r in zip(items, rows, strict=True):
        if r["claim"] != it["claim"]:
            changed["real" if it["rule"] == "real" else "constructed"] += 1
        for pat in ("  ", " .", " ,"):
            assert r["claim"].count(pat) <= it["claim"].count(pat), it["item_id"]
        assert not r["claim"][:1].islower(), it["item_id"]
    assert changed == {"real": 26, "constructed": 6}
    shown = "\n".join(r["claim"] for r in rows)
    for kept in ("(US)", "(EFL)", "(r = 0.35)",
                 "(Carefree, Pragmatic, Uncertain and Worried)", "(aged 11"):
        assert kept in shown, kept


def test_claims_guide_says_citations_are_removed(built):
    guide = "\n".join(_column_a(built.out, "Guide"))
    assert ("Claims are shown without their in-text citations. The paper each claim cites is "
            "given in the cited_title, cited_doi and cited_first_author_year columns.") in guide


# --------------------------------------------------------------------------------------
# Task 3: blind screening sample and workbook
# --------------------------------------------------------------------------------------

SMALL = {"exclude": 4, "include": 2, "judge_demoted": 1, "other_needs_review": 1}
BUILD_SIZES = {"exclude": 8, "include": 4, "judge_demoted": 2, "other_needs_review": 2}


def _screening_record_dict():
    recs = []

    def add(n, prefix, status, reason=""):
        for i in range(n):
            k = len(recs)
            first_excluded = prefix == "Ex" and i == 0
            recs.append({
                "title": "Paper Excluded Zero" if first_excluded else f"Paper {prefix} {i}",
                "abstract": "" if reason == "no_abstract" else f"Abstract text {prefix} {i}.",
                "journal": f"Journal {k}", "year": str(2000 + k),
                "status": status, "needs_review_reason": reason,
                "doi": f"10.1/SENTINELDOI{k}",
                "openalex_id": f"https://openalex.org/SENTINELOA{k}",
                "journal_issn": f"SENTINELISSN{k}", "quote": "SENTINELQUOTE",
                "reason": "SENTINELREASON", "to_confirm": "SENTINELCONFIRM",
                "second_pass": "SENTINELSECOND", "stage": "llm", "outcome": "x",
            })

    add(12, "Ex", "EXCLUDE")
    add(6, "In", "INCLUDE")
    add(4, "Nr", "NEEDS_REVIEW", "not_established")
    add(3, "Na", "NEEDS_REVIEW", "no_abstract")
    return {
        "criteria": {"query": "Does AI feedback help writers?",
                     "inclusion_criteria": ["Empirical study", "About feedback"],
                     "exclusion_criteria": ["Not peer reviewed", "Not in English"]},
        "records": recs,
    }


@pytest.fixture
def record():
    return _screening_record_dict()


@pytest.fixture
def record_path(tmp_path, record):
    path = tmp_path / "screening_record.json"
    path.write_text(json.dumps(record), encoding="utf-8")
    return path


@pytest.fixture
def exclude_keys_path(tmp_path):
    path = tmp_path / "prior_keys.txt"
    path.write_text("title:paper excluded zero\n", encoding="utf-8")
    return path


def _sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _build_screening(record_path, exclude_keys_path, copy="1", sizes=BUILD_SIZES):
    return br.build_blind_screening_workbook(
        record_path, expect_sha256=_sha(record_path), exclude_keys_path=exclude_keys_path,
        seeds=FIXED, copy=copy, max_cell_chars=32_000, sizes=sizes)


def _save_screening(tmp_path, built_tuple, tag):
    wb, sample, ids, workbook_id, _counts = built_tuple
    out, key = tmp_path / f"{tag}.xlsx", tmp_path / f"{tag}.csv"
    wb.save(out)
    br.write_screening_key_csv(key, sample, ids, workbook_id=workbook_id)
    return out, key


@pytest.fixture
def built_screening(tmp_path, record_path, exclude_keys_path):
    out, key = _save_screening(
        tmp_path, _build_screening(record_path, exclude_keys_path), "scr1")
    return types.SimpleNamespace(out=out, key=key, tmp=tmp_path)


def test_stratum_rule():
    nr = {"status": "NEEDS_REVIEW", "needs_review_reason": "not_established"}
    na = {"status": "NEEDS_REVIEW", "needs_review_reason": "no_abstract"}
    assert br.blind_screening_stratum(nr) == "judge_demoted"
    assert br.blind_screening_stratum(na) == "other_needs_review"
    assert br.blind_screening_stratum({"status": "EXCLUDE"}) == "exclude"
    assert br.blind_screening_stratum({"status": "INCLUDE"}) == "include"
    assert br.blind_screening_stratum({"status": "other"}) is None


def test_match_keys_and_normalised_title():
    assert br.normalise_title("  A-Title: Of  X! ") == "a title of x"
    assert br.normalise_title(None) == ""
    keys = br.record_match_keys({"openalex_id": "https://openalex.org/W12", "doi": "10.1/ABC",
                                 "title": "Some Title."})
    assert keys == {"openalex:W12", "doi:10.1/abc", "title:some title"}
    assert br.record_match_keys({"openalex_id": "", "doi": None, "title": ""}) == set()


def test_load_key_lines_reads_one_key_per_line_and_none_means_no_keys(tmp_path):
    txt = tmp_path / "keys.txt"
    txt.write_text("title:alpha\n\n  doi:10.1/b  \n", encoding="utf-8")
    assert br.load_key_lines(txt) == {"title:alpha", "doi:10.1/b"}
    assert br.load_key_lines(None) == set()


def test_sample_sizes_exclusion_and_determinism(record):
    excl = {"title:paper excluded zero"}
    s1, c1 = br.draw_screening_sample(record["records"], sizes=SMALL, seed=5, exclude_keys=excl)
    s2, _ = br.draw_screening_sample(record["records"], sizes=SMALL, seed=5, exclude_keys=excl)
    assert [r["title"] for r in s1] == [r["title"] for r in s2] and len(s1) == 8
    assert "Paper Excluded Zero" not in {r["title"] for r in s1}   # matched by title only
    assert c1["exclude"] == {"drawn": 4, "available": 11}
    assert {r["_stratum"] for r in s1} == set(br.BLIND_SCREENING_STRATA)


def test_short_stratum_taken_whole(record):
    sample, counts = br.draw_screening_sample(
        record["records"], sizes={**SMALL, "judge_demoted": 5}, seed=1, exclude_keys=set())
    short = counts["judge_demoted"]
    assert short["available"] < 5 and short["drawn"] == short["available"]
    assert sum(1 for r in sample if r["_stratum"] == "judge_demoted") == short["drawn"]
    # the other strata keep their requested size, and the short one is not refilled from them
    assert counts["exclude"]["drawn"] == SMALL["exclude"]


def test_wrong_sha_refuses(record_path, exclude_keys_path):
    with pytest.raises(ex.ExportError, match="sha256"):
        br.build_blind_screening_workbook(
            record_path, expect_sha256="0" * 64, exclude_keys_path=exclude_keys_path,
            seeds=FIXED, copy="1", max_cell_chars=32_000, sizes=SMALL)


def test_expect_git_accepts_the_committed_record(record_path, exclude_keys_path, monkeypatch):
    committed = Path(record_path).read_bytes()
    seen = []
    monkeypatch.setattr(br, "_git_show_bytes", lambda rev_path: seen.append(rev_path) or committed)
    wb, sample, _ids, _wid, _counts = br.build_blind_screening_workbook(
        record_path, expect_git="472c465:demo/expected/screening_record.json",
        exclude_keys_path=exclude_keys_path, seeds=FIXED, copy="1", max_cell_chars=32_000,
        sizes=SMALL)
    assert seen == ["472c465:demo/expected/screening_record.json"]
    assert len(sample) == sum(SMALL.values())


def test_expect_git_refuses_a_changed_record_without_printing_a_hash(
        record_path, exclude_keys_path, monkeypatch):
    changed = Path(record_path).read_bytes() + b" "
    monkeypatch.setattr(br, "_git_show_bytes", lambda rev_path: changed)
    with pytest.raises(ex.ExportError) as e:
        br.build_blind_screening_workbook(
            record_path, expect_git="472c465:demo/expected/screening_record.json",
            exclude_keys_path=exclude_keys_path, seeds=FIXED, copy="1", max_cell_chars=32_000,
            sizes=SMALL)
    assert "differs from" in str(e.value)
    assert not re.search(r"[0-9a-f]{32,}", str(e.value))


def test_screening_cli_requires_one_identity_flag(tmp_path, record_path, exclude_keys_path):
    seeds = tmp_path / "seeds.json"
    br.write_seeds_file(seeds, FIXED)
    base = ["screening", "--screening-record", str(record_path),
            "--exclude-keys", str(exclude_keys_path), "--seeds-file", str(seeds),
            "--copy", "1", "--out", str(tmp_path / "o.xlsx"), "--key-out", str(tmp_path / "o.csv")]
    with pytest.raises(SystemExit) as none_given:
        br.main(base)
    assert none_given.value.code == 2
    with pytest.raises(SystemExit) as both_given:
        br.main([*base, "--expect-sha256", _sha(record_path), "--expect-git", "HEAD:x"])
    assert both_given.value.code == 2


def test_screening_build_needs_exactly_one_identity_argument(record_path, exclude_keys_path):
    with pytest.raises(ValueError):
        br.build_blind_screening_workbook(
            record_path, exclude_keys_path=exclude_keys_path, seeds=FIXED, copy="1",
            max_cell_chars=32_000, sizes=SMALL)


def test_screening_header_sheets_and_no_sentinel(built_screening):
    assert tuple(openpyxl.load_workbook(built_screening.out).sheetnames) == br.SCREENING_SHEETS
    for name, text in _zip_texts(built_screening.out):
        assert "SENTINEL" not in text and "not_established" not in text, name
        assert "NEEDS_REVIEW" not in text, name
    ws = openpyxl.load_workbook(built_screening.out)["Screening"]
    assert tuple(c.value for c in ws[1]) == br.BLIND_SCREENING_HEADER


def test_blank_abstract_shows_placeholder(built_screening):
    rows = _sheet_dicts(built_screening.out, "Screening")
    assert len(rows) == 16
    assert sum(r["abstract"] == ex.NO_ABSTRACT_PLACEHOLDER for r in rows) == 2


def test_criterion_dropdown_lists_ids_and_topic(record):
    assert br.criterion_choices(record["criteria"]) == ("I1", "I2", "E1", "E2", "TOPIC")


def test_blind_screening_row_keys_and_venue(record):
    row = br.blind_screening_row(record["records"][0], "BS-001", max_cell_chars=32_000)
    assert tuple(row) == br.BLIND_SCREENING_HEADER
    assert row["venue"] == record["records"][0]["journal"] and row["record_id"] == "BS-001"


def test_screening_layout_guide_and_dropdowns(built_screening):
    wb = openpyxl.load_workbook(built_screening.out)
    ws = wb["Screening"]
    assert ws.freeze_panes == "C2" and 45 <= ws.row_dimensions[2].height <= 409
    header = [c.value for c in ws[1]]
    widths = {"record_id": 10, "title": 45, "abstract": 150, "venue": 28, "year": 8,
              "human_decision": 14, "human_criterion": 14, "human_note": 30}
    assert {k: ws.column_dimensions[_col(header, k)].width for k in widths} == widths
    lists = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert '"include,exclude,unsure"' in lists and '"I1,I2,E1,E2,TOPIC"' in lists
    guide = "\n".join(_column_a(built_screening.out, "Guide"))
    assert "There are 16 rows." in guide and "Does AI feedback help writers?" in guide
    assert "I1. Empirical study" in guide and "E2. Not in English" in guide
    assert br.read_guide_workbook_id(built_screening.out) is not None
    assert tuple(c.value for c in wb["Attestation"][1]) == br.ATTESTATION_HEADER


def test_screening_properties_are_empty(built_screening):
    props = openpyxl.load_workbook(built_screening.out).properties
    for prop in PROPS:
        assert getattr(props, prop) in (None, ""), prop


def test_screening_created_and_modified_are_build_time_only(
        tmp_path, record_path, exclude_keys_path):
    import datetime as dt
    before = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None) - dt.timedelta(minutes=1)
    out, _ = _save_screening(tmp_path, _build_screening(record_path, exclude_keys_path), "ts")
    after = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None) + dt.timedelta(minutes=1)
    props = openpyxl.load_workbook(out).properties
    for prop in ("created", "modified"):
        value = getattr(props, prop)
        assert value is not None and before <= value <= after, prop


def test_copies_share_ids_but_differ_in_row_order(tmp_path, record_path, exclude_keys_path):
    o1, k1 = _save_screening(tmp_path, _build_screening(record_path, exclude_keys_path), "c1")
    o2, k2 = _save_screening(
        tmp_path, _build_screening(record_path, exclude_keys_path, copy="2"), "c2")
    r1, r2 = _sheet_dicts(o1, "Screening"), _sheet_dicts(o2, "Screening")
    assert {r["record_id"]: r["title"] for r in r1} == {r["record_id"]: r["title"] for r in r2}
    assert k1.read_bytes() == k2.read_bytes()
    assert [r["record_id"] for r in r1] != [r["record_id"] for r in r2]


def test_screening_key_has_stratum_and_passes_audit(built_screening):
    key = list(csv.DictReader(built_screening.key.open(encoding="utf-8")))
    assert tuple(key[0]) == ("opaque_id", "stratum", "status", "needs_review_reason",
                             "openalex_id", "doi", "title", "workbook_id")
    assert Counter(r["stratum"] for r in key) == BUILD_SIZES
    assert {r["opaque_id"] for r in key} == {
        r["record_id"] for r in _sheet_dicts(built_screening.out, "Screening")}
    assert br.audit_blind_workbook(
        built_screening.out, kind="screening", key_path=built_screening.key) == []


def test_screening_cli_and_export_dispatch(tmp_path, record_path, exclude_keys_path, monkeypatch):
    monkeypatch.setattr(br, "DEFAULT_SCREENING_SIZES", BUILD_SIZES)
    seeds = tmp_path / "seeds.json"
    br.write_seeds_file(seeds, FIXED)

    def argv(tag, sha=None):
        return ["--screening-record", str(record_path), "--expect-sha256",
                sha or _sha(record_path), "--exclude-keys", str(exclude_keys_path),
                "--seeds-file", str(seeds), "--copy", "1", "--out", str(tmp_path / f"{tag}.xlsx"),
                "--key-out", str(tmp_path / f"{tag}.csv")]

    assert br.main(["screening", *argv("a")]) == 0
    assert ex.main(["--blind-screening", *argv("b")]) == 0
    assert (_sheet_dicts(tmp_path / "a.xlsx", "Screening")
            == _sheet_dicts(tmp_path / "b.xlsx", "Screening"))
    assert br.main(["screening", *argv("c", sha="0" * 64)]) == 1


# --------------------------------------------------------------------------------------
# Manifest and filled-workbook validator
# --------------------------------------------------------------------------------------

BIG_SEEDS = br.BlindSeeds(
    salt="ab" * 16, claims_id_seed=918273645, claims_row_seeds=(912345678, 923456789, 934567890),
    screening_sample_seed=945678901, screening_id_seed=956789012,
    screening_row_seeds=(967890123, 978901234, 989012345))


def test_item_and_record_hashes():
    assert br.item_sha256({"item_id": "a", "claim": "c"}) == hashlib.sha256(
        b"a\x00c").hexdigest()
    assert br.record_sha256({"openalex_id": "https://openalex.org/W9", "doi": "10.1/x",
                             "title": "T"}) == hashlib.sha256(b"openalex:W9").hexdigest()
    assert br.record_sha256({"doi": "10.1/X", "title": "T"}) == hashlib.sha256(
        b"doi:10.1/x").hexdigest()
    assert br.record_sha256({"title": "A  Title!"}) == hashlib.sha256(
        b"title:a title").hexdigest()


def test_manifest_has_commitments_and_never_a_seed(tmp_path, record_path, exclude_keys_path):
    items = [_constructed(i) for i in range(1, 5)]
    items.append(_constructed(5, rule="no_full_text", chunk=False))
    items += [_real(i) for i in range(1, 4)]
    items_path, cache = _write_items(tmp_path, items)
    m = br.build_manifest(
        [(items_path, cache)], record_path, exclude_keys_path, BIG_SEEDS,
        exporter_commit="abc123", sizes=BUILD_SIZES)
    text = json.dumps(m)
    s = BIG_SEEDS
    for v in (s.claims_id_seed, s.screening_sample_seed, s.screening_id_seed,
              *s.claims_row_seeds, *s.screening_row_seeds):
        assert str(v) not in text
    assert s.salt not in text
    assert m["exporter_commit"] == "abc123" and m["generated"]
    assert m["commitment_rule"] == (
        "sha256 of '<seed>:<salt>', seed in decimal; seeds and salt are published with the "
        "results")
    assert m["rubric_sha256"] == hashlib.sha256(
        "\n".join(ex.STATUS_DEFINITIONS_V3).encode()).hexdigest()
    c = m["claims"]
    assert c["n_items"] == 8 and c["n_no_full_text"] == 1
    assert c["item_sha256"] == sorted(br.item_sha256(it) for it in items)
    assert c["id_seed_commitment"] == br.seed_commitment(s.claims_id_seed, s.salt)
    assert c["row_seed_commitments"] == [br.seed_commitment(v, s.salt)
                                         for v in s.claims_row_seeds]
    sc = m["screening"]
    assert sc["record_sha256_file"] == _sha(record_path)
    assert sc["n_records"] == 16 and len(sc["record_sha256"]) == 16
    assert sc["record_sha256"] == sorted(sc["record_sha256"])
    assert sc["strata"] == BUILD_SIZES
    assert sc["n_records_left_out_by_key"] == 1
    assert sc["sample_seed_commitment"] == br.seed_commitment(s.screening_sample_seed, s.salt)
    assert sc["id_seed_commitment"] == br.seed_commitment(s.screening_id_seed, s.salt)
    assert len(sc["row_seed_commitments"]) == 3
    assert set(m) == {"generated", "exporter_commit", "commitment_rule", "rubric_sha256",
                      "claims", "screening"}


def _fill_attestation(wb, **override):
    ws = wb["Attestation"]
    answers = ["BSc, linguistics", "Lecturer", "10", "none", "no", "yes", "yes", "yes", "yes",
               "Pat Example", "45", "2000-01-01"]
    for r, value in enumerate(answers, start=2):
        ws.cell(row=r, column=2, value=value)
    for ref, value in override.items():
        ws.cell(row=int(ref[1:]), column=2).value = value


def _filled_claims(tmp_path, *, verdict="verified", minutes=5, name="filled.xlsx", **att):
    path = tmp_path / name
    wb = openpyxl.Workbook()
    wb.active.title = "Guide"
    ws = wb.create_sheet("Claims")
    ws.append(list(br.BLIND_CLAIMS_HEADER))
    for vals in (
        {"item_id": "BC-001", "human_verdict": verdict, "confidence": "high",
         "human_quote": "q", "human_note": "n", "minutes_spent": minutes},
        {"item_id": "BC-002", "human_verdict": "unsupported", "confidence": "low",
         "human_note": "n", "minutes_spent": 3},
    ):
        vals = {"claim": "c", "stored_text": "stored", "cited_title": "t", "cited_doi": "d",
                "cited_first_author_year": "a", **vals}
        ws.append([vals.get(h, "") for h in br.BLIND_CLAIMS_HEADER])
    br.write_attestation_sheet(wb)
    _fill_attestation(wb, **att)
    wb.save(path)
    return path


def test_validate_flags_case_and_whitespace(tmp_path):
    p = _filled_claims(tmp_path, verdict="Verified ")
    f = br.validate_filled_workbook(p, kind="claims")
    assert len(f) == 1 and "BC-001" in f[0] and "human_verdict" in f[0]


def test_validate_accepts_complete_workbook(tmp_path):
    assert br.validate_filled_workbook(_filled_claims(tmp_path), kind="claims") == []


@pytest.mark.parametrize("minutes", [-1, "x", "abc", 2.5])
def test_validate_flags_bad_minutes(tmp_path, minutes):
    f = br.validate_filled_workbook(_filled_claims(tmp_path, minutes=minutes), kind="claims")
    assert len(f) == 1 and "BC-001" in f[0] and "minutes_spent" in f[0]


@pytest.mark.parametrize("minutes", [None, "", "   "])
def test_validate_accepts_blank_minutes_cell(tmp_path, minutes):
    p = _filled_claims(tmp_path, minutes=minutes)
    assert br.validate_filled_workbook(p, kind="claims") == []


def test_validate_accepts_deleted_minutes_column(tmp_path):
    p = _filled_claims(tmp_path)
    wb = openpyxl.load_workbook(p)
    ws = wb["Claims"]
    ws.delete_cols(list(br.BLIND_CLAIMS_HEADER).index("minutes_spent") + 1)
    wb.save(p)
    assert br.validate_filled_workbook(p, kind="claims") == []


def test_deleted_minutes_column_does_not_hide_other_problems(tmp_path):
    p = _filled_claims(tmp_path, verdict="mostly")
    wb = openpyxl.load_workbook(p)
    wb["Claims"].delete_cols(list(br.BLIND_CLAIMS_HEADER).index("minutes_spent") + 1)
    wb.save(p)
    f = br.validate_filled_workbook(p, kind="claims")
    assert len(f) == 1 and "BC-001" in f[0] and "human_verdict" in f[0]


@pytest.mark.parametrize("kind", ["claims", "screening"])
@pytest.mark.parametrize("blank", [None, "", "   "])
def test_validate_accepts_blank_total_minutes(tmp_path, kind, blank):
    if kind == "claims":
        p = _filled_claims(tmp_path, B12=blank)
    else:
        p = _filled_screening(tmp_path, decision="include", B12=blank)
    assert br.validate_filled_workbook(p, kind=kind) == []


@pytest.mark.parametrize("kind", ["claims", "screening"])
def test_validate_still_flags_wrong_total_minutes(tmp_path, kind):
    if kind == "claims":
        p = _filled_claims(tmp_path, B12="abc")
    else:
        p = _filled_screening(tmp_path, decision="include", B12="abc")
    f = br.validate_filled_workbook(p, kind=kind)
    assert len(f) == 1 and "Total minutes" in f[0]


def test_validate_flags_unknown_verdict_and_confidence(tmp_path):
    p = _filled_claims(tmp_path, verdict="mostly")
    wb = openpyxl.load_workbook(p)
    wb["Claims"]["D3"] = "sure"
    wb.save(p)
    f = br.validate_filled_workbook(p, kind="claims")
    assert len(f) == 2
    assert any("BC-001" in x and "human_verdict" in x for x in f)
    assert any("BC-002" in x and "confidence" in x for x in f)


def test_validate_attestation_rules(tmp_path):
    # B6 is a yes/no answer that is neither yes nor no, B3 an empty text row, B12 a non-integer
    # total; B11 (the name row) is empty while consent in B10 is yes.
    p = _filled_claims(tmp_path, B6="maybe", B3="", B11="", B12="soon")
    f = br.validate_filled_workbook(p, kind="claims")
    assert len(f) == 4
    assert any("Attestation" in x and "maybe" in x for x in f)
    assert any("name" in x for x in f)
    assert any("Total minutes" in x for x in f)
    assert any("Current role" in x for x in f)
    ok = _filled_claims(tmp_path, name="nc.xlsx", B10="no", B11="")
    assert br.validate_filled_workbook(ok, kind="claims") == []


@pytest.mark.parametrize("kind", ["claims", "screening"])
def test_attestation_zero_is_an_answer_and_blank_is_not(tmp_path, kind):
    def build(name, **att):
        if kind == "claims":
            return _filled_claims(tmp_path, name=name, **att)
        return _filled_screening(tmp_path, decision="include", name=name, **att)

    assert br.validate_filled_workbook(build("zero.xlsx", B4=0), kind=kind) == []
    for i, blank in enumerate(("", "   ", None)):
        f = br.validate_filled_workbook(build(f"blank{i}.xlsx", B4=blank), kind=kind)
        assert len(f) == 1 and f[0].endswith("the cell is empty."), f
    assert br.validate_filled_workbook(build("name0.xlsx", B11=0), kind=kind) == []
    f = br.validate_filled_workbook(build("name_sp.xlsx", B11="  "), kind=kind)
    assert f == [f[0]] and "the name is empty although you agreed to be named." in f[0]


def test_criterion_zero_is_not_blank(tmp_path):
    f = br.validate_filled_workbook(
        _filled_screening(tmp_path, criterion=0), kind="screening")
    assert f == []


INTEGRITY_CASES = [
    ("ai_tool_used", "B7", "no", "yes"),
    ("seen_before", "B6", "yes", "no"),
    ("discussed_with_others", "B8", "no", "yes"),
]


@pytest.mark.parametrize(("key", "cell", "flag_value", "ok_value"), INTEGRITY_CASES)
def test_integrity_flags_each_answer(tmp_path, key, cell, flag_value, ok_value):
    flagged = _filled_claims(tmp_path, name="f.xlsx", **{cell: flag_value})
    assert br.validate_filled_workbook(flagged, kind="claims") == []
    flags = br.integrity_flags(flagged)
    assert [k for k, _ in flags] == [key]
    assert flags[0][1].startswith("Integrity flag: Attestation row ")
    ok = _filled_claims(tmp_path, name="o.xlsx", **{cell: ok_value})
    assert br.integrity_flags(ok) == []
    blank = _filled_claims(tmp_path, name="b.xlsx", **{cell: None})
    assert br.integrity_flags(blank) == []
    assert len(br.validate_filled_workbook(blank, kind="claims")) == 1


def test_validate_cli_prints_integrity_warnings_and_keeps_exit_0(tmp_path, capsys):
    p = _filled_claims(tmp_path, B7="no", B6="yes")
    assert br.main(["validate", str(p), "--kind", "claims"]) == 0
    out = capsys.readouterr().out.splitlines()
    assert sum(1 for ln in out if ln.startswith("warning: Integrity flag")) == 2
    assert ("integrity flags: ai_tool_used=1, seen_before=1, discussed_with_others=0"
            in out)


def test_integrity_counts_over_three_workbooks(tmp_path):
    a = _filled_claims(tmp_path, name="a.xlsx", B7="no")
    b = _filled_claims(tmp_path, name="b.xlsx")
    c = _filled_claims(tmp_path, name="c.xlsx", B7="no", B8="no")
    counts = br.integrity_counts([a, b, c])
    assert counts == {"n_workbooks": 3, "ai_tool_used": 2, "seen_before": 0,
                      "discussed_with_others": 1, "flagged_workbooks": [1, 3]}


def _filled_screening(tmp_path, *, decision="exclude", criterion="", name="filled_s.xlsx",
                      **att):
    path = tmp_path / name
    wb = openpyxl.Workbook()
    wb.active.title = "Guide"
    ws = wb.create_sheet("Screening")
    ws.append(list(br.BLIND_SCREENING_HEADER))
    for vals in (
        {"record_id": "BS-001", "human_decision": decision, "human_criterion": criterion,
         "human_note": "n"},
        {"record_id": "BS-002", "human_decision": "include"},
    ):
        vals = {"title": "T", "abstract": "A", "venue": "V", "year": 2020, **vals}
        ws.append([vals.get(h, "") for h in br.BLIND_SCREENING_HEADER])
    br.write_attestation_sheet(wb)
    _fill_attestation(wb, **att)
    wb.save(path)
    return path


def test_validate_requires_criterion_for_exclude(tmp_path):
    f = br.validate_filled_workbook(
        _filled_screening(tmp_path), kind="screening", criteria_choices=("I1", "E1", "TOPIC"))
    assert len(f) == 1 and "BS-001" in f[0] and "human_criterion" in f[0]


def test_validate_screening_accepts_and_flags_criterion(tmp_path):
    choices = ("I1", "E1", "TOPIC")
    assert br.validate_filled_workbook(
        _filled_screening(tmp_path, criterion="E1"), kind="screening",
        criteria_choices=choices) == []
    f = br.validate_filled_workbook(
        _filled_screening(tmp_path, criterion="E9"), kind="screening",
        criteria_choices=choices)
    assert len(f) == 1 and "BS-001" in f[0] and "human_criterion" in f[0]
    f = br.validate_filled_workbook(
        _filled_screening(tmp_path, decision="Exclude", criterion="E1"), kind="screening",
        criteria_choices=choices)
    assert len(f) == 1 and "BS-001" in f[0] and "human_decision" in f[0]


def test_cli_subcommands(tmp_path, record_path, exclude_keys_path, capsys, monkeypatch):
    seeds = tmp_path / "seeds.json"
    br.write_seeds_file(seeds, BIG_SEEDS)
    items = [_constructed(i) for i in range(1, 5)]
    items.append(_constructed(5, rule="no_full_text", chunk=False))
    items += [_real(i) for i in range(1, 5)]
    items_path, cache = _write_items(tmp_path, items)
    base = ["manifest", "--items", str(items_path), "--cache-dir", str(cache),
            "--screening-record", str(record_path), "--exclude-keys",
            str(exclude_keys_path), "--seeds-file", str(seeds), "--out",
            str(tmp_path / "m.json")]
    with pytest.raises(SystemExit):  # --exporter-commit is required
        br.main(base)
    capsys.readouterr()
    # no default record identity: one of --expect-sha256 and --expect-git is required
    with pytest.raises(SystemExit):
        br.main([*base, "--exporter-commit", "abc"])
    capsys.readouterr()
    # a record that is not the expected file is refused, and the message holds no digest
    assert br.main([*base, "--exporter-commit", "abc", "--expect-sha256", "0" * 64]) == 1
    refusal = capsys.readouterr().out
    assert "does not match" in refusal and not re.search(r"[0-9a-f]{32,}", refusal)
    monkeypatch.setattr(br, "DEFAULT_SCREENING_SIZES", BUILD_SIZES)
    assert br.main([*base, "--exporter-commit", "abc", "--expect-sha256",
                    _sha(record_path)]) == 0
    written = (tmp_path / "m.json").read_text(encoding="utf-8")
    assert json.loads(written)["exporter_commit"] == "abc"
    assert str(BIG_SEEDS.claims_id_seed) not in written and BIG_SEEDS.salt not in written
    filled = _filled_claims(tmp_path)
    assert br.main(["validate", str(filled), "--kind", "claims"]) == 0
    bad = _filled_claims(tmp_path, verdict="Verified ", name="bad.xlsx")
    assert br.main(["validate", str(bad), "--kind", "claims"]) == 1
    scr = _filled_screening(tmp_path)
    assert br.main(["validate", str(scr), "--kind", "screening", "--screening-record",
                    str(record_path)]) == 1  # BS-001 is excluded with no criterion


def test_only_the_three_sent_copies_exist_and_no_built_in_identity_remains(
        tmp_path, record_path):
    seeds = tmp_path / "seeds.json"
    br.write_seeds_file(seeds, FIXED)
    base = ["screening", "--screening-record", str(record_path), "--expect-sha256",
            _sha(record_path), "--seeds-file", str(seeds), "--out", str(tmp_path / "o.xlsx"),
            "--key-out", str(tmp_path / "o.csv")]
    with pytest.raises(SystemExit):
        br.main([*base, "--copy", "4"])
    assert not hasattr(br, "V110_SCREENING_RECORD_SHA256")
    assert not hasattr(br, "SENT_COPIES_DIR")


def test_exclude_keys_flag_is_optional_and_described_neutrally(tmp_path, record_path,
                                                               capsys, monkeypatch):
    monkeypatch.setattr(br, "DEFAULT_SCREENING_SIZES", BUILD_SIZES)
    seeds = tmp_path / "seeds.json"
    br.write_seeds_file(seeds, FIXED)
    base = ["screening", "--screening-record", str(record_path), "--expect-sha256",
            _sha(record_path), "--seeds-file", str(seeds), "--copy", "1"]
    assert br.main([*base, "--out", str(tmp_path / "a.xlsx"),
                    "--key-out", str(tmp_path / "a.csv")]) == 0
    titles = {r["title"] for r in csv.DictReader((tmp_path / "a.csv").open(encoding="utf-8"))}
    keys = tmp_path / "keys.txt"
    keys.write_text("title:paper excluded zero\n", encoding="utf-8")
    assert br.main([*base, "--exclude-keys", str(keys), "--out", str(tmp_path / "b.xlsx"),
                    "--key-out", str(tmp_path / "b.csv")]) == 0
    titles_b = {r["title"] for r in csv.DictReader((tmp_path / "b.csv").open(encoding="utf-8"))}
    assert "Paper Excluded Zero" not in titles_b
    assert titles != titles_b or "Paper Excluded Zero" not in titles
    capsys.readouterr()
    with pytest.raises(SystemExit):
        br.main(["screening", "--help"])
    assert "record keys to leave out of the draw" in capsys.readouterr().out.replace("\n", " ")


# --------------------------------------------------------------------------------------
# Workbook usability: time guidance, dropdown enforcement, Sources parts, filters, surnames
# --------------------------------------------------------------------------------------


def test_screening_guide_points_to_total_minutes_and_unsure(built_screening):
    guide = "\n".join(_column_a(built_screening.out, "Guide"))
    question = next(q for q, _ in br.ATTESTATION_QUESTIONS
                    if q.startswith(br._TOTAL_MINUTES_PREFIX))
    assert question == "Total minutes spent on this workbook"
    assert (
        "Time and uncertainty. Expect about 1 to 2 minutes per record and about 15 to 30 minutes "
        "for this workbook. The Screening sheet has no time column. When you have finished, "
        "write the total minutes you spent on this workbook in the Attestation row "
        f'"{question}". If the title and abstract do not let you decide a record, choose '
        "unsure; there is no separate confidence column."
    ) in guide
    paragraphs = _column_a(built_screening.out, "Guide")
    at = next(i for i, p in enumerate(paragraphs) if p.startswith("Time and uncertainty."))
    assert paragraphs[at + 1].startswith("Leave every other cell as it is.")
    assert tuple(br.BLIND_SCREENING_HEADER) == (
        "record_id", "title", "human_decision", "human_criterion", "human_note", "abstract",
        "venue", "year")


def test_every_dropdown_rejects_values_outside_its_list(built, built_screening):
    for path in (built.out, built_screening.out):
        wb = openpyxl.load_workbook(path)
        seen = 0
        for ws in wb.worksheets:
            for dv in ws.data_validations.dataValidation:
                seen += 1
                assert dv.showErrorMessage is True, (path, ws.title)
                assert dv.errorStyle in (None, "stop")
                assert dv.errorTitle == "Not in the list"
                assert dv.error == "Choose a value from the dropdown list."
        assert seen >= 3


def test_source_parts_break_at_whitespace_and_rejoin_exactly():
    text = "word " * 2000
    parts = br.split_source_text(text, 2500)
    assert all(len(p) <= 2500 for p in parts) and "".join(parts) == text
    assert all(p[-1].isspace() for p in parts[:-1])
    odd = ("alpha beta gamma " * 400) + "tail"
    odd_parts = br.split_source_text(odd, 1000)
    assert "".join(odd_parts) == odd and all(len(p) <= 1000 for p in odd_parts)
    assert all(p[-1].isspace() for p in odd_parts[:-1])
    solid = "x" * 5200
    assert [len(p) for p in br.split_source_text(solid, 2500)] == [2500, 2500, 200]
    assert br.split_source_text("", 2500) == []


def test_claims_cli_default_sources_part_is_2500(tmp_path):
    items_path, cache = _write_items(
        tmp_path, [_constructed(1)], caches={"10.1/c1": ["word " * 1200]})
    seeds_path = _seeds(tmp_path)
    assert br.main(["claims", *_argv(tmp_path, items_path, cache, "1", "o", seeds_path)]) == 0
    src = _sheet_dicts(tmp_path / "o.xlsx", "Sources")
    assert len(src) == 3 and all(len(s["text"]) <= 2500 for s in src)
    assert br.BLIND_SOURCES_MAX_CHARS == 2500
    guide = "\n".join(_column_a(tmp_path / "o.xlsx", "Guide"))
    assert ("If a Sources cell holds more text than its row shows, select the cell and read "
            "the text in the formula bar.") in guide


def test_data_sheets_have_autofilters(built, built_screening):
    wb = openpyxl.load_workbook(built.out)
    assert wb["Claims"].auto_filter.ref == f"A1:K{built.n_items + 1}"
    src = wb["Sources"]
    assert src.auto_filter.ref == f"A1:C{src.max_row}"
    scr = openpyxl.load_workbook(built_screening.out)["Screening"]
    assert scr.auto_filter.ref == f"A1:H{scr.max_row}"


def test_first_author_year_keeps_lowercase_particles():
    assert br.first_author_year(["de Vocht", "Laherto"], 2017) == "de Vocht & Laherto, 2017"
    assert (br.first_author_year(["van der Heijden", "Beijaard", "Geldens"], 2018)
            == "van der Heijden et al., 2018")
    assert br.first_author_year(["Jan van der Heijden"], 2018) == "van der Heijden, 2018"
    assert br.first_author_year(["Van Nguyen"], 2020) == "Nguyen, 2020"
    assert br.first_author_year(["Ahmet Cihat Yavuz"], 2020) == "Yavuz, 2020"


# ---- row heights sized from the text (blinding review 2: B1, m1) ----

def test_screening_row_height_follows_the_longest_wrapped_cell():
    width = br._SCREENING_WIDTHS["abstract"]
    short = {"title": "Short title", "abstract": "Short abstract."}
    long_ = {"title": "Short title", "abstract": "w" * 2800}
    need = -(-2800 // width) * br._SCREENING_LINE_HEIGHT
    assert br.screening_row_height(short) < br.screening_row_height(long_)
    assert br.screening_row_height(long_) >= min(need, 409)
    long_title = {"title": "t" * 400, "abstract": "Short abstract."}
    title_need = -(-400 // br._SCREENING_WIDTHS["title"]) * br._SCREENING_LINE_HEIGHT
    assert br.screening_row_height(long_title) >= title_need


def test_screening_row_height_is_capped_and_overflow_is_counted():
    huge = {"title": "T", "abstract": "w" * 20000}
    assert br.screening_row_height(huge) == 409
    rows = [huge, {"title": "T", "abstract": "ok"}, huge]
    assert br.count_overflow_rows(rows) == 2
    assert br.count_overflow_rows([rows[1]]) == 0


def test_workbook_heights_fit_a_2800_character_abstract(tmp_path, exclude_keys_path):
    rec = _screening_record_dict()
    for k, r in enumerate(x for x in rec["records"] if x["status"] == "EXCLUDE"):
        r["abstract"] = "word " * 559 + "words" if k % 2 == 0 else "word " * 3999 + "words"
    path = tmp_path / "rec.json"
    path.write_text(json.dumps(rec), encoding="utf-8")
    out, _key = _save_screening(tmp_path, _build_screening(path, exclude_keys_path), "tall")
    ws = openpyxl.load_workbook(out)["Screening"]
    abstract_col = [c.value for c in ws[1]].index("abstract") + 1
    by_len = {}
    for r in range(2, ws.max_row + 1):
        text = str(ws.cell(r, abstract_col).value or "")
        by_len.setdefault(len(text), []).append(ws.row_dimensions[r].height)
    assert 2800 in by_len and max(by_len) > 6000
    for n, hs in by_len.items():
        want = min(409, max(45, -(-n // 150) * 15))
        assert set(hs) == {want}, (n, hs, want)
    assert set(by_len[2800]) == {285} and set(by_len[max(by_len)]) == {409}
    assert ws.column_dimensions[_col(br.BLIND_SCREENING_HEADER, "abstract")].width >= 147


def test_a_long_abstract_fits_one_row():
    width = br._SCREENING_WIDTHS["abstract"]
    assert br.screening_row_height({"title": "T", "abstract": "w" * 2739}) < 409
    assert width * 27 >= 3969


def test_join_hard_wraps_makes_paragraphs():
    raw = ("The first line of a\nparagraph that was wrapped\nby the PDF.\n\n"
           "Second para-\ngraph, split by hyph-\nenation.\n\n\nThird\tone\nends.")
    out = br.join_hard_wraps(raw)
    assert out == ("The first line of a paragraph that was wrapped by the PDF.\n\n"
                   "Second para-graph, split by hyph-enation.\n\nThird\tone ends.")
    assert "\n" not in br.join_hard_wraps("a\nb\nc")


def test_sources_parts_never_exceed_the_lines_of_a_409_point_row(tmp_path):
    line = "alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu nu xi omicron\n"
    text = (line * 30 + "\n") * 14
    wrapped = br.join_hard_wraps(text)
    assert wrapped.count("\n") < text.count("\n") / 5
    parts = br.split_source_text(
        wrapped, 2500, max_lines=br._SOURCES_MAX_LINES, chars_per_line=br._SOURCES_WIDTH)
    assert "".join(parts) == wrapped
    for p in parts:
        lines = sum(max(1, -(-len(s) // br._SOURCES_WIDTH)) for s in p.split("\n"))
        assert len(p) <= 2500 and lines <= br._SOURCES_MAX_LINES
    assert br._SOURCES_MAX_LINES * br._SOURCES_LINE_HEIGHT <= 409
    many_short = "\n\n".join(["tiny words here"] * 200)
    parts = br.split_source_text(
        many_short, 2500, max_lines=br._SOURCES_MAX_LINES, chars_per_line=br._SOURCES_WIDTH)
    assert "".join(parts) == many_short and len(parts) > 2
    assert all(sum(max(1, -(-len(s) // br._SOURCES_WIDTH)) for s in p.split("\n"))
               <= br._SOURCES_MAX_LINES for p in parts)


def test_blind_sources_display_joins_lines_but_item_text_is_untouched(tmp_path):
    items_path, cache = _write_items(
        tmp_path, [_constructed(1)],
        caches={"10.1/c1": ["Line one of\nthe source text\nruns on.\n\nNew paragraph\nhere."]})
    wb, _ordered, _ids, _wid = br.build_blind_claims_workbook(
        [(items_path, cache)], year_sources=[], seeds=FIXED, copy="1",
        sources_max_chars=2500, max_cell_chars=32_000)
    cells = [str(c.value) for c in wb["Sources"]["C"][1:]]
    shown = "".join(cells)
    assert "Line one of the source text runs on." in shown and "New paragraph here." in shown
    assert "of\nthe" not in shown


# ---- blinding review 3: formulas, claims row heights, carriage returns ----

@pytest.mark.parametrize("lead", ["=", "+", "-", "@"])
def test_free_text_starting_with_formula_characters_is_stored_as_text(tmp_path, lead, capsys):
    text = f"{lead}6.200, p =.013 and more words of the source text."
    item = _constructed(1)
    item["claim"] = f"{lead}6.200 is a claim."
    items_path, cache = _write_items(tmp_path, [item], caches={"10.1/c1": [text]})
    seeds_path = _seeds(tmp_path)
    argv = _argv(tmp_path, items_path, cache, "1", "frm", seeds_path)
    assert br.main(["claims", *argv]) == 0, capsys.readouterr().out
    wb = openpyxl.load_workbook(tmp_path / "frm.xlsx")
    part = wb["Sources"]["C2"]
    assert part.value == text and part.data_type == "s"
    claim = wb["Claims"]["B2"]
    assert claim.value == f"{lead}6.200 is a claim." and claim.data_type == "s"
    for _name, xml in _zip_texts(tmp_path / "frm.xlsx"):
        assert "<f>" not in xml and "<f " not in xml
    assert br._audit(tmp_path / "frm.xlsx", "claims", tmp_path / "frm.csv")[0] == []


@pytest.mark.parametrize("lead", ["=", "+", "-", "@"])
def test_screening_text_starting_with_formula_characters_is_stored_as_text(
        tmp_path, exclude_keys_path, lead):
    rec = _screening_record_dict()
    for r in rec["records"]:
        r["abstract"] = f"{lead}1+1 abstract" if r["abstract"] else ""
        r["title"] = f"{lead}Title {r['title']}"
    path = tmp_path / "rec.json"
    path.write_text(json.dumps(rec), encoding="utf-8")
    out, key = _save_screening(tmp_path, _build_screening(path, exclude_keys_path), "frm")
    ws = openpyxl.load_workbook(out)["Screening"]
    col = [c.value for c in ws[1]].index("abstract") + 1
    shown = [ws.cell(r, col) for r in range(2, ws.max_row + 1)
             if "1+1" in str(ws.cell(r, col).value)]
    assert shown and all(c.value.startswith(f"{lead}1+1") and c.data_type == "s" for c in shown)
    assert br._audit(out, "screening", key)[0] == []


def test_a_real_formula_is_still_flagged(tmp_path):
    path = _clean_claims_wb(tmp_path / "f.xlsx", leak="formula")
    assert any("holds a formula" in f for f in br._audit(path, "claims", None)[0])


def test_claims_row_height_follows_the_longest_wrapped_cell():
    row = {"claim": "x" * 323, "cited_title": "T", "cited_doi": "", "stored_text": "s",
           "cited_first_author_year": "A, 2020"}
    assert br.claims_row_height(row) == 7 * 15
    assert br.claims_row_height({**row, "claim": "short"}) == 45
    assert br.claims_row_height({**row, "claim": "w" * 50000}) == 409
    title = {**row, "claim": "short", "cited_title": "t" * 400}
    assert br.claims_row_height(title) == -(-400 // 35) * 15 == 180
    rows = [{**row, "claim": "w" * 50000}, row, {**row, "claim": "w" * 50000}]
    assert br.count_claims_overflow_rows(rows) == 2


def test_claims_workbook_gives_a_323_character_claim_its_line_count(tmp_path, capsys):
    item = _constructed(1)
    item["claim"] = "word " * 64 + "abc"
    assert len(item["claim"]) == 323
    items_path, cache = _write_items(tmp_path, [item])
    seeds_path = _seeds(tmp_path)
    argv = _argv(tmp_path, items_path, cache, "1", "tall", seeds_path)
    assert br.main(["claims", *argv]) == 0
    assert "rows taller than the 409-point maximum: 0 of 1" in capsys.readouterr().out
    ws = openpyxl.load_workbook(tmp_path / "tall.xlsx")["Claims"]
    assert ws.row_dimensions[2].height == 7 * 15


CR, LF = chr(13), chr(10)


def test_line_breaks_cr_crlf_and_lf_are_treated_alike():
    assert br.join_hard_wraps(CR.join("abc")) == "a b c"
    assert br.join_hard_wraps((CR + LF).join("abc")) == "a b c"
    cr_text = CR.join(["p one", "line", "", "p two"])
    lf_text = LF.join(["p one", "line", "", "p two"])
    assert br.join_hard_wraps(cr_text) == br.join_hard_wraps(lf_text) == "p one line\n\np two"
    assert br.join_hard_wraps((CR + LF).join(["p one", "", "p two"])) == "p one\n\np two"
    assert br.join_hard_wraps("hyph-" + CR + "enation") == "hyph-enation"
    assert br.wrapped_line_count("a" + CR + "b" + CR + LF + "c" + LF + "d", 120) == 4


def test_cr_only_source_text_gives_parts_within_the_line_cap(tmp_path):
    line = "alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu nu xi omicron"
    raw = CR.join([line] * 400)
    items_path, cache = _write_items(
        tmp_path, [_constructed(1)], caches={"10.1/c1": [raw]})
    wb, _o, _i, _w = br.build_blind_claims_workbook(
        [(items_path, cache)], year_sources=[], seeds=FIXED, copy="1",
        sources_max_chars=br.BLIND_SOURCES_MAX_CHARS, max_cell_chars=32_000)
    parts = [str(c.value) for c in wb["Sources"]["C"][1:]]
    assert parts
    for p in parts:
        assert CR not in p
        assert sum(max(1, -(-len(s) // br._SOURCES_WIDTH))
                   for s in p.split(LF)) <= br._SOURCES_MAX_LINES


# ---- Sources display text: searchable characters and sentence-end parts ----

def test_fold_unicode_spaces():
    raw = "negative\u00a0concerns\u2009and\u202fp\u2003q\u200br"
    assert br.fold_unicode_spaces(raw) == "negative concerns and p qr"
    odd = "a" + chr(9) + "b" + chr(10) + "c" + chr(13) + chr(10) + "d"
    assert br.fold_unicode_spaces(odd) == odd


def test_ligature_rejoin_in_split_style_text():
    raw = "the \ufb01 rst \ufb01 ndings o\ufb00 ers di\ufb00 erences a\ufb03  rm"
    assert br.repair_extraction_glyphs(raw) == "the first findings offers differences affirm"


def test_ligature_rejoin_only_in_split_style_text():
    assert br.repair_extraction_glyphs(
        "rounds o\ufb00 the \ufb01nal e\ufb00ect") == "rounds off the final effect"
    assert br.repair_extraction_glyphs("Wol\ufb00, C.") == "Wolff, C."


def test_tt_glyph_repair():
    raw = "a\u0134 itudes, h\u0134 ps, Ma\u0134 ers, A\u0134 itude, Foske\u0134  (Eds.)"
    assert br.repair_extraction_glyphs(raw) == (
        "attitudes, https, Matters, Attitude, Foskett  (Eds.)")


def test_split_prefers_sentence_end_and_rejoins_exactly():
    sentence = "See Fig. 3 for the results of Smith et al. (2010). "
    text = (sentence * 120)[:6000]
    parts = br.split_source_text(text, 2500)
    assert "".join(parts) == text and all(len(p) <= 2500 for p in parts)
    ends = {m.end() for m in br._SENTENCE_END.finditer(text)}
    cut = 0
    for part in parts[:-1]:
        cut += len(part)
        assert cut in ends and part[-1].isspace()
    flat = "word " * 2000
    flat_parts = br.split_source_text(flat, 2500)
    assert "".join(flat_parts) == flat
    assert all(p[-1].isspace() for p in flat_parts[:-1])


def test_quote_copied_from_display_matches_display_text():
    import common

    stored = "Their a\u0134 itudes were\u00a0positive."
    quote = "Their attitudes were positive."
    assert common.quote_is_verbatim(quote, br.display_source_text(stored))


def test_display_source_text_is_the_three_steps():
    stored = "Line a\u0134 itude" + chr(10) + "runs\u00a0on o\ufb00 ers."
    assert br.display_source_text(stored) == br.repair_extraction_glyphs(
        br.join_hard_wraps(br.fold_unicode_spaces(stored)))
    assert "\u00a0" not in br.display_source_text(stored)


def test_sources_sheet_shows_display_text(tmp_path):
    stored = "Their a\u0134 itudes were\u00a0positive."
    items_path, cache = _write_items(tmp_path, [_constructed(1)], caches={"10.1/c1": [stored]})
    wb, _o, _i, _w = br.build_blind_claims_workbook(
        [(items_path, cache)], year_sources=[], seeds=FIXED, copy="1",
        sources_max_chars=2500, max_cell_chars=32_000)
    assert wb["Sources"]["C2"].value == "Their attitudes were positive."


CORPUS_ITEM_FILES = ("hss_test_v3_claims.jsonl", "real_claims_test.jsonl")
CORPUS_CACHES = ("hss_test_v3_fulltext", "real_claims_fulltext")
_corpus_data = Path(os.environ.get("BLIND_FULLTEXT_DATA", ""), "hss_test_v3_fulltext")


def _corpus_texts():
    specs = [(CLAIMS_DIR / f, Path(os.environ["BLIND_FULLTEXT_DATA"]) / c)
             for f, c in zip(CORPUS_ITEM_FILES, CORPUS_CACHES, strict=True)]
    items = ex.read_items(specs)
    caches = ex.build_caches(specs)
    texts = {}
    for it in items:
        doi = it.get("chunk_doi")
        if doi and doi not in texts:
            texts[doi] = ex.sanitize_for_excel(
                ex.full_text_for_doi(doi, caches[it["_cache_dir"]]))
    return texts


@pytest.mark.skipif(not _corpus_data.is_dir(),
                    reason="cached full texts are not in this checkout")
def test_display_normalisation_over_the_86_item_texts():
    texts = _corpus_texts()
    assert len(texts) == 47
    shown = {doi: br.display_source_text(t) for doi, t in texts.items()}
    for doi, t in shown.items():
        assert not br._UNSEARCHABLE.search(t), doi
    counts = {
        ("10.18806/tesl.v36i1.1301", "attitude"): 34,
        ("10.30935/scimath/9511", "negative concerns"): 11,
        ("10.1007/s10734-010-9356-0", "financial"): 27,
    }
    for (doi, term), n in counts.items():
        assert shown[doi].lower().count(term) == n, (doi, term)
        assert texts[doi].lower().count(term) == 0, (doi, term)
    v6n1 = next(t for d, t in shown.items() if d.endswith("v6n1.1168"))
    assert v6n1.lower().count("off the") == 2
    n_non_final = n_mid = 0
    end = re.compile("(?:[.!?][\"'\u201d\u2019)\\]]*[ \t]+|\n\n+)$")
    for doi, t in shown.items():
        wb = openpyxl.Workbook()
        ws = wb.active
        br._write_sources_sheet(
            ws, [{"item_id": doi}], {doi: "BC-001"}, {doi: texts[doi]}, 2500)
        parts = [str(c.value) for c in ws["C"][1:]]
        assert "".join(parts) == t, doi
        for part in parts:
            assert len(part) <= 2500
            assert br.wrapped_line_count(part, br._SOURCES_WIDTH) <= br._SOURCES_MAX_LINES
        for part in parts[:-1]:
            n_non_final += 1
            n_mid += 0 if end.search(part) else 1
    assert n_mid <= 0.01 * n_non_final, (n_mid, n_non_final)


# ---- decision columns next to the claim and the title ----

def _px(width):
    return 7 * width + 5


def test_decision_columns_sit_within_1300_px(built, built_screening):
    assert br.BLIND_CLAIMS_HEADER == (
        "item_id", "claim", "human_verdict", "confidence", "human_quote", "human_note",
        "minutes_spent", "stored_text", "cited_title", "cited_doi", "cited_first_author_year")
    assert br.BLIND_SCREENING_HEADER == (
        "record_id", "title", "human_decision", "human_criterion", "human_note", "abstract",
        "venue", "year")
    for path, sheet, last in ((built.out, "Claims", "minutes_spent"),
                              (built_screening.out, "Screening", "human_note")):
        ws = openpyxl.load_workbook(path)[sheet]
        header = [c.value for c in ws[1]]
        assert ws.freeze_panes == "C2"
        total = sum(_px(ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width)
                    for i in range(header.index(last) + 1))
        assert total <= 1300, (sheet, total)


def test_column_letter_follows_the_header():
    assert br._column_letter(br.BLIND_CLAIMS_HEADER, "human_verdict") == "C"
    assert br._column_letter(br.BLIND_CLAIMS_HEADER, "confidence") == "D"
    assert br._column_letter(br.BLIND_SCREENING_HEADER, "human_decision") == "C"
    assert br._column_letter(br.BLIND_SCREENING_HEADER, "human_criterion") == "D"


def test_built_dropdowns_sit_on_the_decision_columns(built, built_screening):
    ws = openpyxl.load_workbook(built.out)["Claims"]
    got = {dv.formula1: {c for rng in dv.sqref.ranges for c in _cells(rng)}
           for dv in ws.data_validations.dataValidation}
    assert got['"high,medium,low"'] == {f"D{r}" for r in range(2, 22)}
    ws = openpyxl.load_workbook(built_screening.out)["Screening"]
    got = {dv.formula1: {c for rng in dv.sqref.ranges for c in _cells(rng)}
           for dv in ws.data_validations.dataValidation}
    n = ws.max_row
    assert got['"include,exclude,unsure"'] == {f"C{r}" for r in range(2, n + 1)}
    assert got['"I1,I2,E1,E2,TOPIC"'] == {f"D{r}" for r in range(2, n + 1)}


OLD_CLAIMS_HEADER = (
    "item_id", "claim", "cited_title", "cited_doi", "cited_first_author_year", "stored_text",
    "human_verdict", "confidence", "human_quote", "human_note", "minutes_spent")
OLD_SCREENING_HEADER = (
    "record_id", "title", "abstract", "venue", "year", "human_decision", "human_criterion",
    "human_note")


def test_validate_and_scorers_read_columns_by_name_in_any_order(tmp_path):
    # claims: a filled workbook in the earlier column order
    old = tmp_path / "old_claims.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Guide"
    ws = wb.create_sheet("Claims")
    ws.append(list(OLD_CLAIMS_HEADER))
    for i, verdict in enumerate(("verified", "unsupported", "needs_nuance"), start=1):
        vals = {"item_id": f"BC-{i:03d}", "claim": "c", "human_verdict": verdict,
                "confidence": "high", "minutes_spent": 4}
        ws.append([vals.get(h, "") for h in OLD_CLAIMS_HEADER])
    br.write_attestation_sheet(wb)
    _fill_attestation(wb)
    wb.save(old)
    assert br.validate_filled_workbook(old, kind="claims") == []
    import human_review_labels as hrl

    key = tmp_path / "key.csv"
    with key.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["opaque_id", "real_item_id", "set", "construction_category",
                    "expected_label", "workbook_id"])
        for i in range(1, 4):
            w.writerow([f"BC-{i:03d}", f"item-{i}", "real", "", "verified", "w"])
    labels = hrl.build_human_labels((old, old, old), key, sheet_name="Claims")
    assert len(labels["items"]) == 3
    # screening: three filled copies in the earlier column order
    import blind_screening_scores as bss

    skey = tmp_path / "skey.csv"
    with skey.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=br.SCREENING_KEY_HEADER)
        w.writeheader()
        for i in range(1, 4):
            w.writerow({"opaque_id": f"BS-{i:03d}", "stratum": "exclude", "status": "",
                        "needs_review_reason": "", "openalex_id": "", "doi": "",
                        "title": "t", "workbook_id": "0123456789abcdef"})
    paths = []
    for c in range(3):
        path = tmp_path / f"old_s{c}.xlsx"
        wb = openpyxl.Workbook()
        guide = wb.active
        guide.title = "Guide"
        guide["A1"] = br.GUIDE_ID_LINE.format(
            workbook_id="0123456789abcdef", copy="1", export_time="x")
        ws = wb.create_sheet("Screening")
        ws.append(list(OLD_SCREENING_HEADER))
        for i in range(1, 4):
            vals = {"record_id": f"BS-{i:03d}", "title": "t", "human_decision": "exclude",
                    "human_criterion": "E1"}
            ws.append([vals.get(h, "") for h in OLD_SCREENING_HEADER])
        br.write_attestation_sheet(wb)
        _fill_attestation(wb)
        wb.save(path)
        paths.append(path)
        assert br.validate_filled_workbook(path, kind="screening") == []
    assert bss.score_blind_screening(paths, skey)["n_records"] == 3


# ---- Returned-workbook tolerances ----

@pytest.mark.parametrize("answer", ["No", " no ", "NO", "Yes", " yes ", "YES	"])
def test_validate_accepts_yes_no_in_any_case_with_spaces(tmp_path, answer):
    p = _filled_claims(tmp_path, B6=answer, B7=answer, B10=answer, B11="Pat Example")
    assert br.validate_filled_workbook(p, kind="claims") == []


def test_validate_still_flags_a_yes_no_answer_that_is_neither(tmp_path):
    f = br.validate_filled_workbook(_filled_claims(tmp_path, B6="Nope"), kind="claims")
    assert len(f) == 1 and "Attestation row 6" in f[0]


def test_free_text_attestation_answers_are_unchanged(tmp_path):
    p = _filled_claims(tmp_path, B3="   ")
    f = br.validate_filled_workbook(p, kind="claims")
    assert len(f) == 1 and "Current role" in f[0]


def test_integrity_flags_read_yes_no_case_insensitively(tmp_path):
    p = _filled_claims(tmp_path, B7=" No ", B6=" Yes ")
    keys = {k for k, _ in br.integrity_flags(p)}
    assert keys == {"ai_tool_used", "seen_before"}


def _drop_total_minutes_row(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Attestation"]
    row = next(r for r in range(2, ws.max_row + 1)
               if str(ws.cell(row=r, column=1).value).startswith("Total minutes"))
    ws.delete_rows(row)
    wb.save(path)


@pytest.mark.parametrize("kind", ["claims", "screening"])
def test_validate_accepts_workbook_without_total_minutes_question(tmp_path, kind):
    p = (_filled_claims(tmp_path) if kind == "claims"
         else _filled_screening(tmp_path, decision="include"))
    _drop_total_minutes_row(p)
    assert br.validate_filled_workbook(p, kind=kind) == []


def test_removed_total_minutes_row_does_not_hide_other_missing_questions(tmp_path):
    p = _filled_claims(tmp_path)
    _drop_total_minutes_row(p)
    wb = openpyxl.load_workbook(p)
    wb["Attestation"].delete_rows(3)
    wb.save(p)
    f = br.validate_filled_workbook(p, kind="claims")
    assert any("Current role" in x and "missing" in x for x in f)
    assert not any("Total minutes" in x for x in f)


# ---- Identifying the sent copy of a returned workbook ----

def _sent_copy(sent_dir, kind, copy, ids, workbook_id):
    wb = openpyxl.Workbook()
    wb.active.title = "Guide"
    wb.active["A1"] = br.GUIDE_ID_LINE.format(
        workbook_id=workbook_id, copy=copy, export_time="t")
    ws = wb.create_sheet("Claims" if kind == "claims" else "Screening")
    ws.append(["item_id" if kind == "claims" else "record_id", "x"])
    for i in ids:
        ws.append([i, ""])
    wb.save(sent_dir / f"blind-{kind}-copy{copy}.xlsx")


def _returned(path, kind, ids, guide_line=None):
    wb = openpyxl.Workbook()
    wb.active.title = "Guide"
    wb.active["A1"] = guide_line or "Guide text without the id line."
    ws = wb.create_sheet("Claims" if kind == "claims" else "Screening")
    ws.append(["item_id" if kind == "claims" else "record_id", "x"])
    for i in ids:
        ws.append([i, "filled"])
    wb.save(path)
    return path


SENT_IDS = {"1": ("BC-003", "BC-001", "BC-002"), "2": ("BC-001", "BC-002", "BC-003"),
            "3": ("BC-002", "BC-003", "BC-001")}
SENT_WB_IDS = {"1": "1" * 16, "2": "2" * 16, "3": "3" * 16}


@pytest.fixture
def sent_dir(tmp_path):
    d = tmp_path / "sent"
    d.mkdir()
    for kind in ("claims", "screening"):
        for copy in "123":
            _sent_copy(d, kind, copy, SENT_IDS[copy], SENT_WB_IDS[copy])
    return d


def test_resolve_uses_the_guide_line_when_present(tmp_path, sent_dir):
    p = _returned(tmp_path / "r.xlsx", "claims", SENT_IDS["2"],
                  guide_line=br.GUIDE_ID_LINE.format(
                      workbook_id="a" * 16, copy="2", export_time="t"))
    assert br.resolve_workbook_id(p, kind="claims", sent_dir=sent_dir) == (
        "a" * 16, "guide_line")


def test_resolve_uses_the_explicit_copy_argument(tmp_path, sent_dir):
    p = _returned(tmp_path / "r.xlsx", "claims", ("BC-009",))
    assert br.resolve_workbook_id(p, kind="claims", copy="3", sent_dir=sent_dir) == (
        SENT_WB_IDS["3"], "explicit_copy")


@pytest.mark.parametrize("kind", ["claims", "screening"])
@pytest.mark.parametrize("copy", ["1", "2", "3"])
def test_resolve_matches_the_row_order_without_a_copy_argument(tmp_path, sent_dir, kind, copy):
    p = _returned(tmp_path / "r.xlsx", kind, SENT_IDS[copy])
    assert br.resolve_workbook_id(p, kind=kind, sent_dir=sent_dir) == (
        SENT_WB_IDS[copy], "row_order")


def test_resolve_refuses_when_no_sent_copy_matches(tmp_path, sent_dir):
    p = _returned(tmp_path / "r.xlsx", "claims", ("BC-001", "BC-003", "BC-002"))
    with pytest.raises(ValueError, match="no sent copy"):
        br.resolve_workbook_id(p, kind="claims", sent_dir=sent_dir)


def test_resolve_refuses_a_different_number_of_rows(tmp_path, sent_dir):
    p = _returned(tmp_path / "r.xlsx", "claims", SENT_IDS["2"][:2])
    with pytest.raises(ValueError, match="no sent copy"):
        br.resolve_workbook_id(p, kind="claims", sent_dir=sent_dir)


def test_resolve_refuses_when_two_sent_copies_match(tmp_path, sent_dir):
    _sent_copy(sent_dir, "claims", "3", SENT_IDS["2"], SENT_WB_IDS["3"])
    p = _returned(tmp_path / "r.xlsx", "claims", SENT_IDS["2"])
    with pytest.raises(ValueError, match="more than one"):
        br.resolve_workbook_id(p, kind="claims", sent_dir=sent_dir)


def test_resolve_refuses_an_unknown_explicit_copy(tmp_path, sent_dir):
    p = _returned(tmp_path / "r.xlsx", "claims", SENT_IDS["2"])
    with pytest.raises(ValueError, match="copy"):
        br.resolve_workbook_id(p, kind="claims", copy="4", sent_dir=sent_dir)


def test_resolve_never_prints_or_returns_the_id_in_the_method(tmp_path, sent_dir):
    p = _returned(tmp_path / "r.xlsx", "claims", SENT_IDS["1"])
    _, method = br.resolve_workbook_id(p, kind="claims", sent_dir=sent_dir)
    assert method in br.COPY_IDENTIFICATION_METHODS and SENT_WB_IDS["1"] not in method


# ---- Guide text: search method, label boundary, time ----

def _guide_paragraphs(kind, n_rows):
    wb = openpyxl.Workbook()
    if kind == "claims":
        br._write_claims_guide(wb.active, n_rows, "0123456789abcdef", "1", "t")
    else:
        criteria = {"query": "Q?", "inclusion_criteria": ["a"], "exclusion_criteria": ["b"]}
        br._write_screening_guide(wb.active, criteria, n_rows, "0123456789abcdef", "1", "t")
    return [c.value for c in wb.active["A"] if c.value]


def test_time_phrase():
    assert br.time_phrase(58.6) == "about 60 minutes"
    assert br.time_phrase(16.07) == "about 15 minutes"
    assert br.time_phrase(840) == "about 14 hours"
    assert br.time_phrase(150) == "about 2.5 hours"
    assert br.time_phrase(2) == "about 5 minutes"


def test_guide_estimates_match_the_round_constants():
    assert br.FULL_ROUND_CLAIM_ITEMS == 86
    assert round(br.CLAIMS_MINUTES_PER_ITEM * 86 / 60, 1) == br.CLAIMS_HOURS_ESTIMATE
    assert round(br.SCREENING_MINUTES_PER_RECORD * 140 / 60, 1) == br.SCREENING_HOURS_ESTIMATE


CLAIMS_NEW_PARAGRAPHS = (
    "How to work through a row. Read the claim and pick two to four key terms from it, such "
    "as a number, a named group or a distinctive word. Filter the Sources item_id column to "
    "the row's item_id, then search the sheet (Ctrl+F, or Cmd+F on a Mac) for each term. Read "
    "the passage around each hit, then decide. Read the whole paper only when the searches do "
    "not settle the row. A long paper is split over several Sources rows, numbered in the "
    "part column. If a Sources cell holds more text than its row shows, select the cell and "
    "read the text in the formula bar.",
    "A sentence may continue in the next part. Read the next part as well, and when you quote "
    "across two parts, join them with one space.",
    "The Sources text is the stored text of the paper with its line breaks joined. Unusual "
    "spaces, quote marks, special hyphens, minus signs and joined letter pairs in the stored "
    "text were replaced with plain characters, so a typed search finds most words. A word "
    "broken at a line end keeps its hyphen (for example stu-dents). If a search finds "
    "nothing, also search for the first part of the word. Copy quotes from it as shown.",
    "If the paper states a different number, direction or detail for the same thing the claim "
    "describes, the paper contradicts the claim: choose unsupported. Choose needs_nuance for a "
    "detail the paper does not state, or for a claim that overstates the paper or leaves out a "
    "qualification it makes.",
    "A number or finding that the cited paper states counts as support, even when the paper "
    "reports it from another source.",
    "Leave every other cell as it is. When you have finished, fill in the Attestation sheet "
    "(write numbers as whole numbers, such as 0 or 12, and the date as YYYY-MM-DD), save the "
    "workbook and send it back.",
)


def test_claims_guide_search_method_parts_and_boundary():
    full = _guide_paragraphs("claims", 86)
    small = _guide_paragraphs("claims", 6)
    for paragraphs in (full, small):
        for text in CLAIMS_NEW_PARAGRAPHS:
            assert text in paragraphs, text[:40]
        last_def = paragraphs.index(ex.STATUS_DEFINITIONS_V3[-1])
        assert paragraphs[last_def + 1] == CLAIMS_NEW_PARAGRAPHS[3]
        assert "Search the text for the claim's key terms" not in "\n".join(paragraphs)
    time_full = next(p for p in full if p.startswith("Time."))
    assert time_full == (
        "Time. Expect about 10 minutes per row on average and about 14 hours for this "
        "workbook. You do not need to finish in one sitting. Write the minutes for each row "
        "in minutes_spent, and the total for the whole workbook in the Attestation sheet.")
    time_small = next(p for p in small if p.startswith("Time."))
    assert "about 60 minutes for this workbook" in time_small


def test_screening_guide_nits_and_time():
    full = _guide_paragraphs("screening", 140)
    small = _guide_paragraphs("screening", 15)
    include = next(p for p in full if p.startswith("include:"))
    assert include.endswith(" Leave human_criterion blank.")
    exclude = next(p for p in full if p.startswith("exclude:"))
    assert exclude.endswith(" or TOPIC if the record is not about the review question at all.")
    assert not any("If several criteria fail" in p for p in full + small)
    assert (
        'A record whose abstract cell reads "(no abstract)" has no abstract, and an abstract '
        'that is cut short (it stops mid-word or ends with "...") is treated the same way. '
        "Decide from the title and whatever abstract text there is, and choose unsure if that "
        "does not settle it.") in full
    time_full = next(p for p in full if p.startswith("Time and uncertainty."))
    assert time_full.startswith(
        "Time and uncertainty. Expect about 1 to 2 minutes per record and about 2.5 to 4.5 hours "
        "for this workbook. The Screening sheet has no time column.")
    time_small = next(p for p in small if p.startswith("Time and uncertainty."))
    assert "about 15 to 30 minutes for this workbook" in time_small
    assert CLAIMS_NEW_PARAGRAPHS[-1] in full


def test_built_guides_pass_the_audit_with_the_new_text(built, built_screening):
    assert br.audit_blind_workbook(built.out, kind="claims", key_path=built.key) == []
    assert br.audit_blind_workbook(
        built_screening.out, kind="screening", key_path=built_screening.key) == []


# ---- blind-fix3: character tidy, wider guard, guide texts, screening cells ----

TIDY_CASES = [
    ("a  b", "a b"),
    ("a \t b", "a b"),
    ("a\t\tb", "a b"),
    ("x\u2010y", "x-y"),
    ("x\u2011y", "x-y"),
    ("3\u22125", "3-5"),
    ("it\u2019s \u2018this\u2019", "it's 'this'"),
    ("\u201cquoted\u201d", '"quoted"'),
]


@pytest.mark.parametrize(("raw", "want"), TIDY_CASES)
def test_display_source_text_tidies_characters(raw, want):
    assert br.display_source_text(raw) == want


def test_display_source_text_keeps_newlines_and_paragraphs_and_stored_text():
    stored = "One  line\nsecond   line\n\nNew  paragraph"
    assert br.display_source_text(stored) == "One line second line\n\nNew paragraph"
    assert stored == "One  line\nsecond   line\n\nNew  paragraph"


def test_display_source_text_joins_a_line_end_hyphen_of_any_kind():
    assert br.display_source_text("stu\u2010\ndents") == "stu-dents"
    assert br.display_source_text("stu\u2011\ndents") == "stu-dents"


@pytest.mark.parametrize("bad", [
    "a\u2010b", "a\u2011b", "a\u2212b", "a\u2018b", "a\u2019b", "a\u201cb", "a\u201db",
    "a  b", "a \t b"])
def test_audit_flags_untidied_sources_text(tmp_path, bad):
    path = _clean_claims_wb(tmp_path / "c.xlsx")
    wb = openpyxl.load_workbook(path)
    wb["Sources"]["C2"] = bad
    _finish(wb, path, None)
    findings = br.audit_blind_workbook(path, kind="claims")
    assert any("Sources" in f and "typed search" in f for f in findings), findings


@pytest.mark.parametrize("col", ["title", "abstract", "venue"])
@pytest.mark.parametrize("bad", ["a\u00a0b", "a\u200bb", "a\u2010b", "a\u2212b", "a\u2019b",
                                 "a\u201db", "a  b"])
def test_audit_flags_untidied_screening_cells(tmp_path, col, bad):
    path = _clean_screening_wb(tmp_path / "s.xlsx")
    wb = openpyxl.load_workbook(path)
    letter = _col(br.BLIND_SCREENING_HEADER, col)
    wb["Screening"][f"{letter}2"] = bad
    _finish(wb, path, None)
    findings = br.audit_blind_workbook(path, kind="screening")
    assert any("Screening" in f and "typed search" in f for f in findings), findings


def test_audit_accepts_ligature_and_plain_text_in_screening_cells(tmp_path):
    path = _clean_screening_wb(tmp_path / "s.xlsx")
    wb = openpyxl.load_workbook(path)
    wb["Screening"][f"{_col(br.BLIND_SCREENING_HEADER, 'title')}2"] = "o\ufb00 er it's"
    _finish(wb, path, None)
    assert br.audit_blind_workbook(path, kind="screening") == []


def test_tidy_display_characters_for_a_screening_cell():
    raw = "A\u00a0study\u200b of  \u201cx\u201d \u2010 y\u2212z \u2019 \u2018"
    assert br.tidy_display_characters(raw) == "A study of \"x\" - y-z ' '"
    assert br.tidy_display_characters("o\ufb00er a\u0134 e") == "o\ufb00er a\u0134 e"


def test_screening_cells_show_tidied_text_and_record_is_untouched(tmp_path, record):
    record["records"][0]["title"] = "Learners\u2019  \u201cvoices\u201d\u00a0matter"
    record["records"][0]["abstract"] = "Line one.\nA\u2011B  and\u200b C \u2212 D"
    record["records"][0]["journal"] = "Journal\u2019s  Name"
    path = tmp_path / "rec.json"
    path.write_text(json.dumps(record), encoding="utf-8")
    before = _sha(path)
    row = br.blind_screening_row(record["records"][0], "BS-001", max_cell_chars=32_000)
    assert row["title"] == "Learners' \"voices\" matter"
    assert row["abstract"] == "Line one.\nA-B and C - D"
    assert row["venue"] == "Journal's Name"
    assert _sha(path) == before


def test_blind_claims_guide_sentence_about_the_tidy_is_true():
    for n in (86, 6):
        text = _guide_paragraphs("claims", n)
        assert not any("so that a search finds the words" in p for p in text)
        assert any("so a typed search finds most words" in p for p in text)
        assert any("special hyphens, minus signs and joined letter pairs" in p for p in text)
        assert not any("quote marks, dashes" in p for p in text)


def test_screening_guide_topic_and_first_failing_criterion_sentence():
    for n in (140, 15):
        paragraphs = _guide_paragraphs("screening", n)
        sentence = (
            "Use TOPIC when the record is not about the review question at all. Otherwise "
            "record the first failing criterion in list order; for unsure, record the first "
            "criterion in list order that the title and abstract cannot settle.")
        assert sentence in paragraphs


def test_screening_hours_constant_is_the_guide_upper_bound():
    assert br.SCREENING_HOURS_ESTIMATE == 4.5
    full = _guide_paragraphs("screening", 140)
    assert any("about 2.5 to 4.5 hours" in p for p in full)


SCREENING_RECORD = os.environ.get("BLIND_SCREENING_RECORD", "")


@pytest.mark.skipif(not Path(SCREENING_RECORD).is_file(),
                    reason="no full-size screening record is named in BLIND_SCREENING_RECORD")
def test_full_size_screening_build_has_no_guard_hits(tmp_path):
    expected = hashlib.sha256(Path(SCREENING_RECORD).read_bytes()).hexdigest()
    wb, sample, ids, workbook_id, _counts = br.build_blind_screening_workbook(
        Path(SCREENING_RECORD), expect_sha256=expected, seeds=FIXED, copy="1",
        max_cell_chars=32_000)
    assert len(sample) == 140
    out = tmp_path / "full.xlsx"
    wb.save(out)
    findings = br.audit_blind_workbook(out, kind="screening")
    assert [f for f in findings if "typed search" in f] == []
