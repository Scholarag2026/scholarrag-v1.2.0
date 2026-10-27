"""`claims/human_review_labels.py`: the majority-label rule, the two inter-rater agreement
statistics, and the v6-vs-human-label scorer, all on synthetic fixtures -- the real
workbooks are gitignored and not part of any checkout, so nothing here depends on them.
"""

from __future__ import annotations

import json
import re

import human_review_labels as hrl
import pytest

# --------------------------------------------------------------------------------------
# Majority rule
# --------------------------------------------------------------------------------------


def test_majority_label_two_of_three_agree():
    assert hrl.majority_label(["verified", "verified", "needs_nuance"]) == "verified"


def test_majority_label_unanimous():
    assert hrl.majority_label(["unsupported", "unsupported", "unsupported"]) == "unsupported"


def test_majority_label_three_way_split_is_no_majority():
    assert hrl.majority_label(["verified", "needs_nuance", "unsupported"]) is None


# --------------------------------------------------------------------------------------
# Inter-rater agreement
# --------------------------------------------------------------------------------------


def test_fleiss_kappa_is_one_when_every_rater_agrees_on_every_item():
    table = [[3, 0], [0, 3], [3, 0]]  # 3 items, 2 categories, unanimous each time
    assert hrl.fleiss_kappa(table) == 1.0


def test_fleiss_kappa_hand_computed_example():
    # 4 items, 3 raters, 2 categories: two unanimous A, one unanimous B, one 2-1 split.
    # P_i = [1, 1, 1, 1/3], Pbar = 5/6; p_A = 8/12, p_B = 4/12, Pe = (2/3)**2 + (1/3)**2 =
    # 5/9; kappa = (5/6 - 5/9) / (1 - 5/9) = (5/18) / (4/9) = 0.625.
    table = [[3, 0], [3, 0], [0, 3], [2, 1]]
    assert hrl.fleiss_kappa(table) == pytest.approx(0.625, abs=1e-9)


def test_fleiss_kappa_none_when_fewer_than_two_raters():
    assert hrl.fleiss_kappa([[1, 0]]) is None


def test_pairwise_agreement_reports_mean_percent_and_kappa():
    cols = {1: ["a", "a", "b"], 2: ["a", "a", "b"], 3: ["a", "b", "b"]}
    out = hrl.pairwise_agreement(cols)
    assert out["mean_pairwise_percent_agreement"] > 0.6
    assert len(out["pairwise_percent_agreement"]) == 3


# --------------------------------------------------------------------------------------
# Scoring v6 against the human majority label
# --------------------------------------------------------------------------------------


def _write_jsonl(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        for r in rows:
            fh.write(json.dumps(r) + "\n")


def test_score_v6_against_labels_computes_accuracy_kappa_and_falsely_verified(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(hrl, "V6_DIR", tmp_path)
    human_labels = {
        "items": [
            {"item_id": "c1", "set": "constructed", "majority_label": "verified"},
            {"item_id": "c2", "set": "constructed", "majority_label": "unsupported"},
            {"item_id": "c3", "set": "constructed", "majority_label": "needs_nuance"},
            {"item_id": "r1", "set": "real", "majority_label": "verified"},
        ],
    }
    _write_jsonl(tmp_path / "hss_runA.jsonl", [
        {"item_id": "c1", "predicted_status": "verified"},
        {"item_id": "c2", "predicted_status": "verified"},  # falsely verified
        {"item_id": "c3", "predicted_status": "needs_nuance"},
    ])
    _write_jsonl(tmp_path / "hss_runB.jsonl", [
        {"item_id": "c1", "predicted_status": "verified"},
        {"item_id": "c2", "predicted_status": "unsupported"},
        {"item_id": "c3", "predicted_status": "needs_nuance"},
    ])
    _write_jsonl(tmp_path / "real_runA.jsonl", [{"item_id": "r1", "predicted_status": "verified"}])
    _write_jsonl(tmp_path / "real_runB.jsonl", [{"item_id": "r1", "predicted_status": "verified"}])

    scores = hrl.score_v6_against_labels(human_labels)

    run_a = scores["by_set"]["constructed"]["run_A"]
    assert run_a["n"] == 3
    assert run_a["n_correct"] == 2
    assert run_a["accuracy"] == pytest.approx(2 / 3, abs=1e-9)
    assert run_a["falsely_verified"] == 1

    run_b = scores["by_set"]["constructed"]["run_B"]
    assert run_b["n_correct"] == 3
    assert run_b["accuracy"] == 1.0
    assert run_b["falsely_verified"] == 0

    real_a = scores["by_set"]["real"]["run_A"]
    assert real_a["n"] == 1
    assert real_a["accuracy"] == 1.0
    assert real_a["verified_precision"] == 1.0
    assert real_a["verified_recall"] == 1.0


# --------------------------------------------------------------------------------------
# CLI: no hard-coded workbook paths, a no-argument run refuses
# --------------------------------------------------------------------------------------


def test_main_refuses_with_no_arguments(capsys):
    rc = hrl.main([])
    assert rc != 0
    assert "ERROR" in capsys.readouterr().out


def test_main_refuses_with_fewer_than_three_workbooks(tmp_path, capsys):
    rc = hrl.main(["--workbook", str(tmp_path / "r1.xlsx"), "--key", str(tmp_path / "k.csv")])
    assert rc != 0
    assert "ERROR" in capsys.readouterr().out


def test_score_v6_against_labels_excludes_no_majority_items(tmp_path, monkeypatch):
    monkeypatch.setattr(hrl, "V6_DIR", tmp_path)
    human_labels = {"items": [
        {"item_id": "c1", "set": "constructed", "majority_label": None},
        {"item_id": "c2", "set": "constructed", "majority_label": "verified"},
    ]}
    _write_jsonl(tmp_path / "hss_runA.jsonl", [
        {"item_id": "c1", "predicted_status": "unsupported"},
        {"item_id": "c2", "predicted_status": "verified"},
    ])
    _write_jsonl(tmp_path / "hss_runB.jsonl", [
        {"item_id": "c1", "predicted_status": "unsupported"},
        {"item_id": "c2", "predicted_status": "verified"},
    ])
    _write_jsonl(tmp_path / "real_runA.jsonl", [])
    _write_jsonl(tmp_path / "real_runB.jsonl", [])

    scores = hrl.score_v6_against_labels(human_labels)
    run_a = scores["by_set"]["constructed"]["run_A"]
    assert run_a["n"] == 1  # the no_majority item is excluded, not scored as wrong
    assert run_a["accuracy"] == 1.0


# --------------------------------------------------------------------------------------
# Blind tag: workbooks with a Claims sheet, joined on the opaque id
# --------------------------------------------------------------------------------------

BLIND_ID = "0123456789abcdef"
KEY_ROWS = [
    ("BC-001", "hss-verbatim-01", "constructed", "verbatim", "verified"),
    ("BC-002", "real-test-01", "real", "", "unsupported"),
    ("BC-003", "hss-paraphrase-01", "constructed", "paraphrase", "needs_nuance"),
]


def _write_key(path, workbook_id=BLIND_ID):
    import csv

    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["opaque_id", "real_item_id", "set", "construction_category",
                    "expected_label", "workbook_id"])
        for oid, real, set_name, cat, exp in KEY_ROWS:
            w.writerow([oid, real, set_name, cat, exp, workbook_id])


def _write_workbook(path, order, verdicts, *, sheet="Claims", workbook_id=BLIND_ID):
    import openpyxl

    wb = openpyxl.Workbook()
    guide = wb.active
    guide.title = "Guide"
    guide.append([f"Workbook id: {workbook_id}. Copy: 1. Exported x."])
    ws = wb.create_sheet(sheet)
    ws.append(["item_id", "human_verdict"])
    for oid in order:
        ws.append([oid, verdicts[oid]])
    wb.save(path)


def _three_copies(tmp_path, verdicts_by_copy, *, sheet="Claims", orders=None):
    orders = orders or [
        ["BC-001", "BC-002", "BC-003"],
        ["BC-003", "BC-001", "BC-002"],
        ["BC-002", "BC-003", "BC-001"],
    ]
    paths = []
    for n, (order, verdicts) in enumerate(zip(orders, verdicts_by_copy, strict=True), start=1):
        p = tmp_path / f"copy{n}.xlsx"
        _write_workbook(p, order, verdicts, sheet=sheet)
        paths.append(p)
    k = tmp_path / "key.csv"
    _write_key(k)
    return paths, k


def _verdict_sets():
    return [
        {"BC-001": "verified", "BC-002": "unsupported", "BC-003": "verified"},
        {"BC-001": "verified", "BC-002": "unsupported", "BC-003": "unsupported"},
        {"BC-001": "verified", "BC-002": "verified", "BC-003": "needs_nuance"},
    ]


def test_blind_tag_joins_on_id_across_reordered_copies(tmp_path):
    (a, b, c), k = _three_copies(tmp_path, _verdict_sets())
    labels = hrl.build_human_labels((a, b, c), k, sheet_name="Claims")
    assert {i["item_id"]: i["majority_label"] for i in labels["items"]} == {
        "hss-verbatim-01": "verified", "real-test-01": "unsupported",
        "hss-paraphrase-01": None}


def test_blind_tag_normalises_case_and_whitespace(tmp_path):
    sets = _verdict_sets()
    sets[0]["BC-001"] = "Verified "
    sets[1]["BC-001"] = "verified"
    sets[2]["BC-001"] = "VERIFIED"
    (a, b, c), k = _three_copies(tmp_path, sets)
    items = hrl.build_human_labels((a, b, c), k, sheet_name="Claims")["items"]
    item = next(i for i in items if i["item_id"] == "hss-verbatim-01")
    assert item["majority_label"] == "verified" and item["unanimous"] is True


def _prepare_main(tmp_path, monkeypatch):
    import summarize

    monkeypatch.setattr(hrl, "V6_DIR", tmp_path)
    monkeypatch.setattr(summarize, "render_markdown", lambda s: "# E2\n\nBASE\n")
    (tmp_path / "summary.json").write_text(
        json.dumps({"generated": "x", "human_review": {"label_of_record": True}}),
        encoding="utf-8")
    for run in ("A", "B"):
        _write_jsonl(tmp_path / f"hss_run{run}.jsonl", [
            {"item_id": "hss-verbatim-01", "predicted_status": "verified"},
            {"item_id": "hss-paraphrase-01", "predicted_status": "needs_nuance"},
        ])
        _write_jsonl(tmp_path / f"real_run{run}.jsonl", [
            {"item_id": "real-test-01", "predicted_status": "unsupported"},
        ])
    return _three_copies(tmp_path, _verdict_sets())


def test_blind_tag_writes_blind_files_only(tmp_path, monkeypatch):
    (a, b, c), k = _prepare_main(tmp_path, monkeypatch)
    argv = ["--tag", "blind", "--workbook", str(a), "--workbook", str(b),
            "--workbook", str(c), "--key", str(k)]
    assert hrl.main(argv) == 0
    assert (tmp_path / "human_labels_blind.json").exists()
    assert (tmp_path / "human_review_scores_blind.json").exists()
    assert {x.name for x in tmp_path.glob("human_*.json")} == {
        "human_labels_blind.json", "human_review_scores_blind.json"}
    scores = json.loads((tmp_path / "human_review_scores_blind.json").read_text(
        encoding="utf-8"))
    assert scores["label_of_record"] == (
        "blind labeller majority (results/v6/human_labels_blind.json)")
    s = json.loads((tmp_path / "summary.json").read_text(encoding="utf-8"))
    assert s["human_review_blind"]["label_of_record"] is True
    assert s["human_review_blind"]["exposure"] == "blind"
    assert "human_review" not in s
    md = (tmp_path / "summary.md").read_text(encoding="utf-8")
    assert "## Table E2-j blind labeller majority, label of record" in md
    assert re.findall(r"^## Table (E2-\w)", md, re.M) == ["E2-j"]
    assert "visible" not in md.lower()
    assert "E2-g" not in md and "E2-d" not in md and "model-authored" not in md


def test_blind_outputs_carry_no_timestamp(tmp_path, monkeypatch):
    (a, b, c), k = _prepare_main(tmp_path, monkeypatch)
    argv = ["--tag", "blind", "--workbook", str(a), "--workbook", str(b),
            "--workbook", str(c), "--key", str(k)]
    assert hrl.main(argv) == 0
    for name in ("human_labels_blind.json", "human_review_scores_blind.json"):
        text = (tmp_path / name).read_text(encoding="utf-8")
        assert '"generated"' not in text
        assert not re.search(r"20\d\d-\d\d-\d\d", text)
    blind_block = json.loads(
        (tmp_path / "summary.json").read_text(encoding="utf-8"))["human_review_blind"]
    assert "generated" not in blind_block
    assert "generated" not in blind_block["scores"]
    assert not re.search(r"20\d\d-\d\d-\d\d", json.dumps(blind_block))


def test_blind_tag_refuses_workbook_id_mismatch(tmp_path):
    (a, b, c), k = _three_copies(tmp_path, _verdict_sets())
    _write_workbook(b, ["BC-001", "BC-002", "BC-003"], _verdict_sets()[1],
                    workbook_id="ffffffffffffffff")
    with pytest.raises(ValueError, match=b.name):
        hrl.check_workbook_ids([a, b, c], k)
    hrl.check_workbook_ids([a, c], k)


def test_tag_accepts_only_blind():
    assert hrl.parse_args(["--workbook", "a", "--key", "k"]).tag == "blind"
    with pytest.raises(SystemExit):
        hrl.parse_args(["--tag", "other", "--workbook", "a", "--key", "k"])


def test_blind_labels_carry_integrity_counts(tmp_path, monkeypatch):
    (a, b, c), k = _prepare_main(tmp_path, monkeypatch)
    argv = ["--tag", "blind", "--workbook", str(a), "--workbook", str(b),
            "--workbook", str(c), "--key", str(k)]
    assert hrl.main(argv) == 0
    labels = json.loads((tmp_path / "human_labels_blind.json").read_text(encoding="utf-8"))
    flags = labels["integrity_flags"]
    assert flags["n_workbooks"] == 3 and flags["flagged_workbooks"] == []
    assert flags["ai_tool_used"] == flags["seen_before"] == flags["discussed_with_others"] == 0


def test_agreement_per_population(tmp_path):
    (a, b, c), k = _three_copies(tmp_path, _verdict_sets())
    agr = hrl.build_human_labels((a, b, c), k, sheet_name="Claims")["inter_rater_agreement"]
    pops = agr["by_population"]
    assert set(pops) == {"constructed", "real"}
    con, real = pops["constructed"], pops["real"]
    assert (con["n_items"], con["n_no_majority"]) == (2, 1)
    assert (real["n_items"], real["n_no_majority"]) == (1, 0)
    assert con["mean_pairwise_percent_agreement"] == pytest.approx(0.5)
    assert real["mean_pairwise_percent_agreement"] == pytest.approx(1 / 3)
    assert "fleiss_kappa" in con and "fleiss_kappa" in real
    # with_withheld is reported beside the judged-only figures, for every population
    assert con["with_withheld"]["n_items"] == 2
    assert real["with_withheld"]["n_no_majority"] == 0


# --------------------------------------------------------------------------------------
# Copy identification when the returned Guide lacks the id line
# --------------------------------------------------------------------------------------

_ORDERS = [
    ["BC-001", "BC-002", "BC-003"],
    ["BC-003", "BC-001", "BC-002"],
    ["BC-002", "BC-003", "BC-001"],
]


def _sent_dir(tmp_path):
    d = tmp_path / "sent"
    d.mkdir()
    for n, order in enumerate(_ORDERS, start=1):
        _write_workbook(d / f"blind-claims-copy{n}.xlsx", order, dict.fromkeys(order, ""))
    return d


def _strip_guide_line(path):
    import openpyxl

    wb = openpyxl.load_workbook(path)
    wb["Guide"]["A1"] = "Guide text only."
    wb.save(path)


def test_check_workbook_ids_matches_by_row_order_when_the_guide_line_is_missing(tmp_path):
    (a, b, c), k = _three_copies(tmp_path, _verdict_sets(), orders=_ORDERS)
    for p in (a, b, c):
        _strip_guide_line(p)
    methods = hrl.check_workbook_ids([a, b, c], k, sent_dir=_sent_dir(tmp_path))
    assert methods == ["row_order", "row_order", "row_order"]


def test_check_workbook_ids_reports_the_guide_line_method(tmp_path):
    (a, b, c), k = _three_copies(tmp_path, _verdict_sets())
    assert hrl.check_workbook_ids([a, b, c], k) == ["guide_line"] * 3


def test_check_workbook_ids_uses_explicit_copies(tmp_path):
    (a, b, c), k = _three_copies(tmp_path, _verdict_sets(), orders=_ORDERS)
    for p in (a, b, c):
        _strip_guide_line(p)
    methods = hrl.check_workbook_ids(
        [a, b, c], k, copies=["1", "2", "3"], sent_dir=_sent_dir(tmp_path))
    assert methods == ["explicit_copy"] * 3


def test_check_workbook_ids_refuses_when_no_sent_copy_matches(tmp_path):
    (a, b, c), k = _three_copies(
        tmp_path, _verdict_sets(),
        orders=[["BC-003", "BC-002", "BC-001"]] + _ORDERS[1:])
    _strip_guide_line(a)
    with pytest.raises(ValueError, match=a.name):
        hrl.check_workbook_ids([a, b, c], k, sent_dir=_sent_dir(tmp_path))


def test_blind_main_records_the_identification_method(tmp_path, monkeypatch):
    (a, b, c), k = _prepare_main(tmp_path, monkeypatch)
    for p in (a, c):
        _strip_guide_line(p)
    sent = _sent_dir(tmp_path)
    argv = ["--tag", "blind", "--workbook", str(a), "--workbook", str(b),
            "--workbook", str(c), "--key", str(k), "--sent-dir", str(sent)]
    assert hrl.main(argv) == 0
    labels = json.loads((tmp_path / "human_labels_blind.json").read_text(encoding="utf-8"))
    assert labels["copy_identification"] == ["row_order", "guide_line", "row_order"]
    assert BLIND_ID not in json.dumps(labels["copy_identification"])


def test_blind_main_takes_a_copy_argument_per_workbook(tmp_path, monkeypatch):
    (a, b, c), k = _prepare_main(tmp_path, monkeypatch)
    for p in (a, b, c):
        _strip_guide_line(p)
    argv = ["--tag", "blind", "--workbook", str(a), "--workbook", str(b),
            "--workbook", str(c), "--key", str(k), "--sent-dir", str(_sent_dir(tmp_path)),
            "--copy", "1", "--copy", "2", "--copy", "3"]
    assert hrl.main(argv) == 0
    labels = json.loads((tmp_path / "human_labels_blind.json").read_text(encoding="utf-8"))
    assert labels["copy_identification"] == ["explicit_copy"] * 3
