"""Tests for the blind screening scorer (Task 7)."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import blind_review as br
import blind_screening_scores as bss
import openpyxl
import pytest

GOOD_ID = "0123456789abcdef"


def _write_workbook(path: Path, decisions: dict[str, str], *, workbook_id: str = GOOD_ID) -> None:
    wb = openpyxl.Workbook()
    guide = wb.active
    guide.title = "Guide"
    guide["A1"] = "Some guide text."
    guide["A2"] = br.GUIDE_ID_LINE.format(
        workbook_id=workbook_id, copy="1", export_time="2026-10-23T00:00:00Z")
    ws = wb.create_sheet("Screening")
    ws.append(list(br.BLIND_SCREENING_HEADER))
    for rid, dec in decisions.items():
        row = {name: "" for name in br.BLIND_SCREENING_HEADER}
        row["record_id"] = rid
        row["title"] = "t"
        row["human_decision"] = dec
        ws.append([row[name] for name in br.BLIND_SCREENING_HEADER])
    br.write_attestation_sheet(wb)
    wb.save(path)


def _write_key(path: Path, strata: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=br.SCREENING_KEY_HEADER)
        w.writeheader()
        for i, stratum in enumerate(strata, start=1):
            w.writerow({"opaque_id": f"BS-{i:03d}", "stratum": stratum, "status": "",
                        "needs_review_reason": "", "openalex_id": f"W{i}", "doi": "",
                        "title": "t", "workbook_id": GOOD_ID})


def _setup(tmp_path, strata, decisions, *, ids=(GOOD_ID,) * 3, drop_from_copy3=()):
    key = tmp_path / "key.csv"
    _write_key(key, strata)
    paths = []
    for c in range(3):
        d = {f"BS-{i + 1:03d}": decisions[i][c] for i in range(len(strata))}
        if c == 2:
            for rid in drop_from_copy3:
                d.pop(rid)
        p = tmp_path / f"copy{c + 1}.xlsx"
        # reverse row order in copy 2 to prove joins use the id
        if c == 1:
            d = dict(reversed(list(d.items())))
        _write_workbook(p, d, workbook_id=ids[c])
        paths.append(p)
    return paths, key


def _score(tmp_path, strata, decisions):
    paths, key = _setup(tmp_path, strata, decisions)
    return bss.score_blind_screening(paths, key)


def test_zero_misses_gives_wilson_upper_bound_0_0458(tmp_path):
    s = _score(tmp_path, ["exclude"] * 80, [("exclude",) * 3] * 80)
    assert s["exclusions"]["n"] == 80 and s["exclusions"]["n_majority_include"] == 0
    assert s["exclusions"]["miss_rate_wilson_95"][1] == pytest.approx(0.0458, abs=1e-4)


def test_three_way_split_is_no_majority_and_left_out_of_rates(tmp_path):
    s = _score(tmp_path, ["exclude"] * 3,
               [("include", "exclude", "unsure"), ("exclude",) * 3, ("exclude",) * 3])
    assert s["n_no_majority"] == 1 and s["by_stratum"]["exclude"]["no_majority"] == 1
    assert s["exclusions"]["n"] == 2 and s["by_stratum"]["exclude"]["n_any_include"] == 1


def test_judge_demoted_upheld_counts_exclude_and_unsure(tmp_path):
    s = _score(tmp_path, ["judge_demoted"] * 4,
               [("exclude",) * 3, ("unsure",) * 3, ("include",) * 3,
                ("include", "exclude", "unsure")])
    jd = s["judge_demoted"]
    assert (jd["n"], jd["n_upheld"], jd["n_majority_include"]) == (3, 2, 1)
    assert jd["rate"] == pytest.approx(2 / 3)


def test_include_agreement(tmp_path):
    s = _score(tmp_path, ["include"] * 3,
               [("include",) * 3, ("include", "include", "exclude"), ("exclude",) * 3])
    assert (s["include_agreement"]["n"], s["include_agreement"]["n_majority_include"]) == (3, 2)


def test_case_and_whitespace_normalised(tmp_path):
    s = _score(tmp_path, ["exclude"], [("Exclude ", "EXCLUDE", "exclude")])
    assert s["by_stratum"]["exclude"]["exclude"] == 1 and s["n_no_majority"] == 0


def test_workbook_id_mismatch_raises(tmp_path):
    paths, key = _setup(tmp_path, ["exclude"], [("exclude",) * 3],
                        ids=(GOOD_ID, "ffffffffffffffff", GOOD_ID))
    with pytest.raises(ValueError, match="workbook"):
        bss.score_blind_screening(paths, key)


def test_copies_with_different_ids_raise(tmp_path):
    paths, key = _setup(tmp_path, ["exclude"] * 2, [("exclude",) * 3] * 2,
                        drop_from_copy3=("BS-002",))
    with pytest.raises(ValueError, match="ids"):
        bss.score_blind_screening(paths, key)


def test_unfilled_decision_raises(tmp_path):
    paths, key = _setup(tmp_path, ["exclude"], [("exclude", "", "exclude")])
    with pytest.raises(ValueError, match="decision"):
        bss.score_blind_screening(paths, key)


def test_inter_rater_and_per_labeller(tmp_path):
    s = _score(tmp_path, ["exclude", "include"],
               [("exclude",) * 3, ("include", "include", "exclude")])
    assert s["inter_rater"]["mean_pairwise_percent_agreement"] == pytest.approx(2 / 3)
    assert s["inter_rater"]["fleiss_kappa"] is not None
    assert s["per_labeller"]["1"] == {"include": 1, "exclude": 1, "unsure": 0}
    assert s["per_labeller"]["3"] == {"include": 0, "exclude": 2, "unsure": 0}


def test_render_markdown_and_cli(tmp_path, capsys):
    paths, key = _setup(tmp_path, ["exclude", "include"],
                        [("exclude",) * 3, ("include",) * 3])
    out = tmp_path / "out"
    rc = bss.main(["--workbook", str(paths[0]), "--workbook", str(paths[1]),
                   "--workbook", str(paths[2]), "--key", str(key), "--out-dir", str(out),
                   "--run-id", "demo-run-X"])
    assert rc == 0
    md = (out / "blind_screening_scores.md").read_text(encoding="utf-8")
    assert md.startswith("## Blind screening sample, 2 records of the delivered run demo-run-X")
    assert json.loads((out / "blind_screening_scores.json").read_text())["n_records"] == 2


def test_scores_carry_integrity_counts(tmp_path):
    s = _score(tmp_path, ["exclude"] * 2, [("exclude",) * 3] * 2)
    assert s["integrity_flags"] == {
        "n_workbooks": 3, "ai_tool_used": 0, "seen_before": 0, "discussed_with_others": 0,
        "flagged_workbooks": []}


def test_ai_tool_answer_no_is_flagged_in_scores_and_in_a_warning_line(tmp_path, capsys):
    paths, key = _setup(tmp_path, ["exclude", "include"],
                        [("exclude",) * 3, ("include",) * 3])
    wb = openpyxl.load_workbook(paths[1])
    att = wb["Attestation"]
    row = next(r for r in range(2, att.max_row + 1)
               if str(att.cell(row=r, column=1).value).startswith("I used no AI tool"))
    att.cell(row=row, column=2).value = "no"
    wb.save(paths[1])
    scores = bss.score_blind_screening(paths, key)
    assert scores["integrity_flags"]["ai_tool_used"] == 1
    assert scores["integrity_flags"]["flagged_workbooks"] == [2]
    rc = bss.main(["--workbook", str(paths[0]), "--workbook", str(paths[1]),
                   "--workbook", str(paths[2]), "--key", str(key),
                   "--out-dir", str(tmp_path / "out"), "--run-id", "r"])
    assert rc == 0
    warnings = [ln for ln in capsys.readouterr().out.splitlines() if ln.startswith("warning:")]
    assert len(warnings) == 1


def test_scores_heading_uses_run_id(tmp_path):
    s = _score(tmp_path, ["exclude"] * 3, [("exclude",) * 3] * 3)
    md = bss.render_markdown(s, run_id="20261106-abc")
    assert md.startswith("## Blind screening sample, 3 records of the delivered run 20261106-abc")
    assert ("v1." + "1.0") not in md and "140" not in md
    other = tmp_path / "other"
    other.mkdir()
    paths, key = _setup(other, ["exclude"], [("exclude",) * 3])
    with pytest.raises(SystemExit):  # --run-id is required
        bss.main(["--workbook", str(paths[0]), "--workbook", str(paths[1]),
                  "--workbook", str(paths[2]), "--key", str(key)])


def test_scores_recompute_from_anonymised_file(tmp_path, capsys):
    strata = ["exclude", "exclude", "include", "judge_demoted"]
    decisions = [("exclude",) * 3, ("include", "exclude", "exclude"),
                 ("include", "include", "unsure"), ("include", "exclude", "unsure")]
    paths, key = _setup(tmp_path, strata, decisions)
    out = tmp_path / "out"
    args = ["--workbook", str(paths[0]), "--workbook", str(paths[1]),
            "--workbook", str(paths[2]), "--key", str(key), "--out-dir", str(out),
            "--run-id", "r"]
    assert bss.main(args) == 0
    records = out / "blind_screening_records.csv"
    text = records.read_text(encoding="utf-8")
    header = text.splitlines()[0].split(",")
    assert header == ["record_id", "openalex_id", "stratum", "system_status",
                      "decision_1", "decision_2", "decision_3"]
    assert text.splitlines()[1].startswith("BS-001,W1,exclude,")
    original = json.loads((out / "blind_screening_scores.json").read_text(encoding="utf-8"))
    again = bss.score_from_records(records)
    original.pop("integrity_flags")
    original.pop("copy_identification")
    assert again == original
    assert bss.main(["--check", str(records), "--out-dir", str(out)]) == 0
    records.write_text(text.replace("W1", "W999"), encoding="utf-8")
    assert bss.main(["--check", str(records), "--out-dir", str(out)]) == 0
    records.write_text(text.replace("include", "exclude", 1), encoding="utf-8")
    assert bss.main(["--check", str(records), "--out-dir", str(out)]) == 1


def _sent_screening(tmp_path, orders):
    d = tmp_path / "sent"
    d.mkdir(exist_ok=True)
    for n, order in enumerate(orders, start=1):
        _write_workbook(d / f"blind-screening-copy{n}.xlsx", dict.fromkeys(order, ""))
    return d


def _strip_guide_line(path):
    wb = openpyxl.load_workbook(path)
    wb["Guide"]["A2"] = None
    wb.save(path)


def test_scorer_identifies_copies_by_row_order_without_the_guide_line(tmp_path):
    orders = [["BS-001", "BS-002", "BS-003"], ["BS-003", "BS-001", "BS-002"],
              ["BS-002", "BS-003", "BS-001"]]
    key = tmp_path / "key.csv"
    _write_key(key, ["exclude"] * 3)
    paths = []
    for n, order in enumerate(orders, start=1):
        p = tmp_path / f"r{n}.xlsx"
        _write_workbook(p, dict.fromkeys(order, "exclude"))
        _strip_guide_line(p)
        paths.append(p)
    sent = _sent_screening(tmp_path, orders)
    scores = bss.score_blind_screening(paths, key, sent_dir=sent)
    assert scores["n_records"] == 3
    assert scores["copy_identification"] == ["row_order"] * 3
    scores = bss.score_blind_screening(paths, key, copies=["1", "2", "3"], sent_dir=sent)
    assert scores["copy_identification"] == ["explicit_copy"] * 3


def test_scorer_refuses_when_no_sent_copy_matches_the_row_order(tmp_path):
    key = tmp_path / "key.csv"
    _write_key(key, ["exclude"] * 2)
    p = tmp_path / "r.xlsx"
    _write_workbook(p, {"BS-001": "exclude", "BS-002": "exclude"})
    _strip_guide_line(p)
    sent = _sent_screening(tmp_path, [["BS-002", "BS-001"]] * 3)
    with pytest.raises(ValueError):
        bss.score_blind_screening([p, p, p], key, sent_dir=sent)


def test_check_ignores_copy_identification(tmp_path):
    paths, key = _setup(tmp_path, ["exclude"], [("exclude",) * 3])
    scores = bss.score_blind_screening(paths, key)
    assert scores["copy_identification"] == ["guide_line"] * 3
    out = tmp_path / "out"
    out.mkdir()
    (out / "blind_screening_scores.json").write_text(json.dumps(scores), encoding="utf-8")
    bss.write_records_file(out / bss.RECORDS_FILE, {"BS-001": "exclude"}, {},
                           [{"BS-001": "exclude"}] * 3)
    assert bss._check(out / bss.RECORDS_FILE, out) == 0
